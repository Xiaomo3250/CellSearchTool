import os, sys
import re
import json
import threading
import socket
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse, parse_qs
import pandas as pd

# 数据库模块
import db
# KML 导出模块
import kml_export

try:
    from openpyxl import load_workbook
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

WORKSPACE = os.path.dirname(os.path.abspath(__file__))
UPLOAD_DIR = os.path.join(WORKSPACE, "uploaded_files")
os.makedirs(UPLOAD_DIR, exist_ok=True)

__version__ = "1.1.0"

# 导入进度状态
import_status = {
    "running": False, "total_files": 0, "current_file": "",
    "current_index": 0, "current_sheet": "", "imported_records": 0,
    "done": False, "error": "",
}


# ============ 字段映射规则 ============
MAPPING_RULES = [
    {
        "carrier": "电信", "tech": "5G",
        "sheet_keywords": ["5G汇总工参", "5G工参", "汇总工参", "NR", "5G"],
        "exclude_keywords": ["删除"],
        "field_map": {
            "运营商": "中国电信", "技术制式": "5G NR",
            "设备商": ["厂家"],
            "基站名": ["基站名称", "网元名称", "基站/楼盘名称", "站址"],
            "基站ID": ["gNodeB标识", "gNodeB ID"],
            "小区名": ["NR小区名称", "小区名称"],
            "小区ID": ["小区ID", "小区本地ID", "NR小区标识", "小区标识"],
            "PCI": ["物理小区标识", "PCI"],
            "下行频点": ["下行频点", "SSB绝对信道号"],
            "下倾角": ["下倾角", "机械下倾角", "电子下倾角"],
            "挂高": ["挂高", "天线挂高", "站高"],
            "方位角": ["方位角"],
            "经度": ["经度"],
            "纬度": ["纬度", "维度"],
            "频段": ["频带", "频段"],
            "共享": ["是否共享", "共享方"],
            "TAC": ["跟踪区码", "跟踪区域码", "TAC"],
        },
    },
    {
        "carrier": "电信", "tech": "4G",
        "sheet_keywords": ["电信工参", "LTE工参", "L1.1室外宏站", "L2.1室内", "室外宏站", "室内"],
        "exclude_keywords": ["删除", "Sheet1"],
        "field_map": {
            "运营商": "中国电信", "技术制式": "4G LTE",
            "设备商": ["设备厂家", "厂家"],
            "基站名": ["EnodebName", "基站名称", "站址"],
            "基站ID": ["EnodebID", "基站ID"],
            "小区名": ["CellName", "小区名称"],
            "小区ID": ["CELLID", "本地CellID", "小区ID"],
            "PCI": ["PCI", "物理小区标识"],
            "下行频点": ["下行中心频点号", "EARFCN", "下行频点", "频点"],
            "下倾角": ["总下倾角", "Downtilt", "下倾角", "机械下倾角"],
            "挂高": ["天线挂高", "GroudHeight", "站高", "挂高"],
            "方位角": ["方位角", "Azimuth"],
            "经度": ["经度", "Longitude"],
            "纬度": ["纬度", "维度", "Latitude"],
            "频段": ["网络类型", "频段", "频带"],
            "共享": ["是否共享", "共享方"],
            "TAC": ["TAC", "跟踪区码", "TAL"],
        },
    },
    {
        "carrier": "联通", "tech": "4G",
        "sheet_keywords": None,
        "exclude_keywords": ["删除", "Sheet1"],
        "field_map": {
            "运营商": "中国联通", "技术制式": "4G LTE",
            "设备商": ["设备厂家", "厂家"],
            "基站名": ["基站名称", "eNodeBName", "站址"],
            "基站ID": ["基站ID", "eNodeBID", "基站标识"],
            "小区名": ["小区名称", "CellName"],
            "小区ID": ["小区ID", "CELLID", "CellID", "本地小区标识"],
            "PCI": ["物理小区标识", "PCI"],
            "下行频点": ["频点", "EARFCN", "下行频点"],
            "下倾角": ["下倾角", "Downtilt", "机械下倾角"],
            "挂高": ["站高", "天线挂高", "GroudHeight", "挂高"],
            "方位角": ["方位角", "Azimuth"],
            "经度": ["经度", "Longitude"],
            "纬度": ["维度", "纬度", "Latitude"],
            "频段": ["频段", "频带"],
            "共享": ["是否共享", "共享方"],
            "TAC": ["跟踪区码", "TAC"],
        },
    },
    {
        "carrier": "联通", "tech": "5G",
        "sheet_keywords": None,
        "exclude_keywords": ["删除", "行政区划"],
        "field_map": {
            "运营商": "中国联通", "技术制式": "5G NR",
            "设备商": ["厂家"],
            "基站名": ["基站名称", "网元名称"],
            "基站ID": ["基站标识", "gNodeB ID"],
            "小区名": ["小区名称"],
            "小区ID": ["NR小区标识", "小区标识"],
            "PCI": ["PCI", "物理小区标识"],
            "下行频点": ["SSB频点", "下行频点"],
            "下倾角": ["电子下倾角", "机械下倾角"],
            "挂高": ["挂高"],
            "方位角": ["方位角"],
            "经度": ["经度"],
            "纬度": ["纬度"],
            "频段": ["频带"],
            "共享": ["是否共享"],
            "TAC": ["跟踪区码", "TAC"],
        },
    },
]


# ============ 工具函数 ============
FILTER_FIELDS = ["制式", "TAC", "基站ID", "小区ID", "PCI", "频点", "基站名", "小区名"]

def _parse_filters(params):
    """从查询参数中提取多字段过滤条件"""
    filters = {}
    for key in FILTER_FIELDS:
        val = params.get(key, [""])[0].strip()
        if val:
            filters[key] = val
    return filters if filters else None


def safe_str(val):
    if val is None: return ""
    s = str(val).strip()
    return "" if s in ("nan", "NaT", "None", "nat") else s


def normalize_share(raw_val, carrier=""):
    """将 Excel 的各种共享列值标准化为 '共享' / '非共享'"""
    v = safe_str(raw_val)
    # "是" / "共享" / "电信" / "Y" → 共享
    if v in ("是", "共享", "Y", "Yes", "电信"):
        return "共享"
    # "否" / "未共享" / "N" / ""(仅电信) → 非共享
    if v in ("否", "未共享", "N", "No") or (v == "" and carrier == "电信"):
        return "非共享"
    # 保留原值（如 "自建"、"共享站"）
    return v if v else ""


def detect_carrier(filename):
    if "电信" in filename: return "电信"
    if "联通" in filename: return "联通"
    if "移动" in filename: return "移动"
    return None


def detect_tech(filename):
    if any(k in filename for k in ("5G", "NR", "5g")): return "5G"
    if any(k in filename for k in ("4G", "LTE", "lte")): return "4G"
    return None


def get_matching_rule(filename):
    carrier = detect_carrier(filename)
    tech = detect_tech(filename)
    for rule in MAPPING_RULES:
        if rule["carrier"] == carrier and rule["tech"] == tech:
            return rule
    return None


def resolve_rule_with_fallback(filepath, filename):
    """获取匹配规则，文件名失败时用首行数据（最可靠）、完整路径、sheet名辅助检测"""
    rule = get_matching_rule(filename)
    if rule is not None:
        return rule

    tech_from_name = detect_tech(filename)
    sheet_names = _read_sheet_names(filepath)

    # 1) 优先读首行数据用命名检测（最可靠，比路径/文件名准确）
    if sheet_names and sheet_names[0]:
        try:
            df_sample = pd.read_excel(filepath, sheet_name=sheet_names[0], nrows=5)
            if len(df_sample) > 0:
                first_row = df_sample.iloc[0].to_dict()
                cell_name = safe_str(
                    first_row.get("NR小区名称") or first_row.get("小区名称") or
                    first_row.get("小区名") or first_row.get("CellName") or ""
                )
                station_name = safe_str(
                    first_row.get("基站名称") or first_row.get("网元名称") or
                    first_row.get("基站名") or first_row.get("EnodebName") or ""
                )
                carrier, _ = detect_carrier_from_names(cell_name, station_name)
                if carrier:
                    c = "电信" if "电信" in carrier else "联通" if "联通" in carrier else None
                    t = tech_from_name or "4G"
                    for r in MAPPING_RULES:
                        if r["carrier"] == c and r["tech"] == t:
                            return r
        except Exception:
            pass

    # 2) 回退：完整路径检测（如 ...\昌吉联通\...）
    carrier_from_path = detect_carrier(filepath.replace("\\", "/"))
    if carrier_from_path and tech_from_name:
        for r in MAPPING_RULES:
            if r["carrier"] == carrier_from_path and r["tech"] == tech_from_name:
                return r

    # 3) 回退：遍历所有规则按 sheet_keywords 匹配 sheet 名
    for r in MAPPING_RULES:
        sk = r.get("sheet_keywords")
        if sk is None:
            continue
        exclude = r.get("exclude_keywords", [])
        for sn in sheet_names:
            if any(kw in sn for kw in exclude):
                continue
            if any(kw in sn for kw in sk):
                if tech_from_name and r["tech"] == tech_from_name:
                    return r
                elif rule is None:
                    rule = r
    if rule:
        return rule

    # 3) 最后兜底：如果有 sheet_keywords=None 的规则（如联通4G/5G），直接用
    for r in MAPPING_RULES:
        if r.get("sheet_keywords") is None and tech_from_name and r["tech"] == tech_from_name:
            return r

    return None


def _read_sheet_names(filepath):
    """快速读取 sheet 名列表"""
    if HAS_OPENPYXL and filepath.lower().endswith((".xlsx", ".xlsm")):
        try:
            wb = load_workbook(filepath, read_only=True)
            names = wb.sheetnames
            wb.close()
            return names
        except Exception:
            pass
    try:
        xls = pd.ExcelFile(filepath)
        names = xls.sheet_names
        xls.close()
        return names
    except Exception:
        return []


# ============ 基于命名的运营商检测 ============
UNICOM_PREFIXES = {"CJCJS","CJFKS","CJQTX","CJMLX","CJHTB","CJMNS","CJJMS","CJWJQ","CJFCH","CJWCW","CJXHN","WLMDQ","WLXSQ","WL_CJ"}
COUNTY_CODE_RE = re.compile(r'^[A-Z]{2,4}\d?$')
END_SEGMENT_RE = re.compile(r'^\d{1,2}$|^.+[EC]$|^.+[EC]-\d+$')

def detect_carrier_from_names(cell_name, station_name):
    """从小区名/基站名判定运营商。返回 (carrier, share_type) 或 (None, None)

    只判定运营商归属，不判定共享状态（共享由 Excel 列决定）。

    规则优先级：
    1. 含 (LTGX) → 电信共享给联通的站（运营商=电信，共享由Excel决定）
    2. 2段 + 联通区县前缀 → 联通自建
    3. 5-7段 + CJ开头 + 区县码格式 + 末段合法 → 电信（共享状态由Excel判定）
    """
    names = [n for n in (cell_name, station_name) if n and n.strip()]
    for name in names:
        if "(LTGX)" in name:
            return ("中国电信", None)  # 电信建、已共享给联通，共享由Excel决定

    for name in names:
        parts = name.split("_")
        n = len(parts)
        first = parts[0] if n > 0 else ""
        last = parts[-1] if n > 0 else ""

        if n == 2 and any(first.startswith(p) for p in UNICOM_PREFIXES):
            return ("中国联通", "自建")

        # 电信命名模式：CJ_区县码_... — 仅判运营商，共享由Excel决定
        if n in (5, 6, 7) and first == "CJ" and COUNTY_CODE_RE.match(parts[1]) and END_SEGMENT_RE.match(last):
            return ("中国电信", None)  # None 表示由Excel列决定共享

    return (None, None)


def get_relevant_sheets(xls, rule):
    """根据规则筛选推荐的 sheet（接收 pd.ExcelFile 对象）"""
    return _filter_sheets(xls.sheet_names, rule)


def get_relevant_sheets_by_names(sheet_names, rule):
    """根据规则筛选推荐的 sheet（接收 sheet 名列表）"""
    return _filter_sheets(sheet_names, rule)


def _filter_sheets(all_names, rule):
    exclude = rule.get("exclude_keywords", [])
    keywords = rule.get("sheet_keywords")
    if keywords is None:
        return [s for s in all_names if not any(kw in s for kw in exclude)]
    matched = []
    for sn in all_names:
        if any(kw in sn for kw in exclude): continue
        if any(kw in sn for kw in keywords):
            matched.append(sn)
    if not matched:
        matched = [s for s in all_names if not any(kw in s for kw in exclude)]
    return matched


def resolve_field(row, field_candidates):
    if isinstance(field_candidates, str): return field_candidates
    for col_name in field_candidates:
        s = safe_str(row.get(col_name))
        if s: return s
    return ""


# ============ 数据库初始化（启动时调用）============
def init_database():
    """初始化 SQLite 数据库，兼容从旧 cache.json 迁移数据"""
    db.init_db()
    # 兼容旧版 cache.json 迁移
    old_cache = os.path.join(WORKSPACE, "cache.json")
    if os.path.exists(old_cache):
        try:
            with open(old_cache, "r", encoding="utf-8") as f:
                data = json.load(f)
            old_records = data.get("records", [])
            old_file_sources = data.get("file_sources", {})
            if old_records:
                # 补全 TAC 字段
                for r in old_records:
                    if "TAC" not in r:
                        r["TAC"] = ""
                    if "_raw" not in r:
                        r["_raw"] = "{}"
                db.insert_records(old_records)
                print(f"[迁移] 从 cache.json 迁移了 {len(old_records)} 条记录")
            for fname, info in old_file_sources.items():
                db.upsert_file_source(
                    fname, info.get("carrier", ""), info.get("tech", ""),
                    info.get("count", 0), info.get("sheets", [])
                )
            # 迁移成功后重命名旧文件
            os.rename(old_cache, old_cache + ".backup")
            print("[迁移] cache.json 已备份为 cache.json.backup")
        except Exception as e:
            print(f"[迁移] 失败: {e}")


# ============ 预览（只解析sheet列表，不导入）============
def preview_file(filepath):
    """返回文件的 sheet 预览信息，不导入数据"""
    filename = os.path.basename(filepath)
    rule = resolve_rule_with_fallback(filepath, filename)
    if rule is None:
        return {"filename": filename, "error": "无法识别运营商/技术制式（文件名需含'电信'或'联通'）", "sheets": []}
    try:
        # 用 openpyxl read_only 模式：一次打开拿所有 sheet 名和 max_row，不解析内容（~1.5s vs 原来 30s+）
        if HAS_OPENPYXL and filepath.lower().endswith((".xlsx", ".xlsm")):
            wb = load_workbook(filepath, read_only=True)
            all_sheets = wb.sheetnames
            # 在同一个 wb 上取每个 sheet 的 max_row（read_only 模式下非常快）
            row_counts = {}
            for sn in all_sheets:
                try:
                    mr = wb[sn].max_row
                    # 过滤空 sheet 或格式化填充行（max_row=1048576 是 Excel 空行上限）
                    row_counts[sn] = mr if 1 <= mr < 500000 else 0
                except Exception:
                    row_counts[sn] = 0
            wb.close()
            sheets_info = []
            # 用 sheet 名列表判定推荐 sheet（不需要 pd.ExcelFile）
            recommended = get_relevant_sheets_by_names(all_sheets, rule)
            for sn in all_sheets:
                sheets_info.append({
                    "name": sn,
                    "rows": row_counts[sn],
                    "recommended": sn in recommended,
                })
        else:
            xls = pd.ExcelFile(filepath)
            all_sheets = xls.sheet_names
            recommended = get_relevant_sheets(xls, rule)
            sheets_info = []
            for sn in all_sheets:
                try:
                    row_count_est = len(pd.read_excel(filepath, sheet_name=sn, usecols=[0]))
                except Exception:
                    row_count_est = 0
                sheets_info.append({
                    "name": sn,
                    "rows": row_count_est,
                    "recommended": sn in recommended,
                })
            xls.close()
        return {
            "filename": filename,
            "carrier": rule["carrier"],
            "tech": rule["tech"],
            "sheets": sheets_info,
        }
    except Exception as e:
        return {"filename": filename, "error": str(e), "sheets": []}


# ============ 导入（指定sheet列表，全量保留原始列）============
def import_file_sheets(filepath, sheet_names, status_update=None):
    """导入指定 sheet，全量保留原始列到 _raw，同时标准化 17 个核心字段"""
    filename = os.path.basename(filepath)
    rule = resolve_rule_with_fallback(filepath, filename)
    if rule is None:
        return 0

    imported_count = 0
    for sn in sheet_names:
        try:
            if status_update:
                status_update(sheet=sn)
            df = pd.read_excel(filepath, sheet_name=sn)
            if df.empty:
                continue
            df = df.fillna("")
            new_records = []
            for _, row in df.iterrows():
                row_dict = row.to_dict()
                # 将所有原始列转为字符串键值对存入 _raw
                raw_dict = {}
                for k, v in row_dict.items():
                    raw_dict[str(k)] = str(v) if v is not None else ""

                rec = {"_文件名": filename, "_工作表": sn, "_raw": raw_dict}
                # 标准化字段映射
                for out_field, candidates in rule["field_map"].items():
                    rec[out_field] = resolve_field(row_dict, candidates)

                # 跳过重复表头行
                cell_name_val = rec.get("小区名", "")
                lng_val = rec.get("经度", "")
                all_candidates = [c for cands in rule["field_map"].values()
                                  if isinstance(cands, list) for c in cands]
                if cell_name_val in all_candidates or lng_val in all_candidates:
                    continue
                # 跳过经纬度明显无效的行
                try:
                    lng_f = float(lng_val)
                    if not (60 <= lng_f <= 140):
                        continue
                except (ValueError, TypeError):
                    if lng_val:
                        continue

                # 运营商判定：优先用"承建方"列，回退到命名检测
                carrier_detected = None
                share_detected = None
                contractor = safe_str(row_dict.get("承建方", ""))
                if "联通" in contractor:
                    carrier_detected = "中国联通"
                elif "电信" in contractor:
                    carrier_detected = "中国电信"
                elif "移动" in contractor:
                    carrier_detected = "中国移动"

                if carrier_detected is None:
                    carrier_detected, share_detected = detect_carrier_from_names(
                        rec.get("小区名", ""), rec.get("基站名", "")
                    )
                if carrier_detected:
                    rec["运营商"] = carrier_detected
                    if share_detected is not None:
                        rec["共享"] = share_detected
                    else:
                        # 通过承建方检测到的运营商，用实际运营商做 normalize
                        carrier_key = "电信" if "电信" in carrier_detected else "联通" if "联通" in carrier_detected else ""
                        rec["共享"] = normalize_share(rec.get("共享", ""), carrier_key)
                else:
                    rec["运营商"] = rule.get("field_map", {}).get("运营商", "")
                    rec["共享"] = normalize_share(rec.get("共享", ""), rule.get("carrier", ""))

                new_records.append(rec)
                imported_count += 1

            if new_records:
                # 批量写入 SQLite
                count = db.insert_records(new_records)
                # 统计该文件多数运营商
                carrier_counts = {}
                for r in new_records:
                    c = r.get("运营商", "")
                    carrier_counts[c] = carrier_counts.get(c, 0) + 1
                majority_carrier = max(carrier_counts, key=carrier_counts.get) if carrier_counts else rule["carrier"]
                # 更新文件来源
                existing_fs = db.get_file_sources().get(filename, {})
                existing_sheets = existing_fs.get("sheets", [])
                if sn not in existing_sheets:
                    existing_sheets.append(sn)
                existing_count = existing_fs.get("count", 0) + count
                db.upsert_file_source(filename, majority_carrier, rule["tech"], existing_count, existing_sheets)

                if status_update:
                    import_status["imported_records"] = db.get_total_count()
            print(f"  [Sheet] {sn}: {len(new_records)} 条")
        except Exception as e:
            print(f"  [Error] sheet {sn}: {e}")

    return imported_count


# ============ 删除指定文件/sheet ============
def remove_sheets(filename, sheet_names=None):
    """
    sheet_names=None  → 删除该文件所有记录
    sheet_names=[...] → 只删除指定 sheet 的记录
    """
    if sheet_names is None:
        db.delete_by_filename(filename)
        db.remove_file_source(filename)
    else:
        db.delete_by_filename_and_sheets(filename, sheet_names)
        # 更新 file_sources
        existing_fs = db.get_file_sources().get(filename, {})
        existing_sheets = existing_fs.get("sheets", [])
        for sn in sheet_names:
            if sn in existing_sheets:
                existing_sheets.remove(sn)
        if existing_sheets:
            db.upsert_file_source(
                filename, existing_fs.get("carrier", ""), existing_fs.get("tech", ""),
                existing_fs.get("count", 0) - len(sheet_names) if existing_fs.get("count", 0) > 0 else 0,
                existing_sheets
            )
        else:
            db.remove_file_source(filename)


# ============ 搜索 ============
def search_records(query, page=1, per_page=50, filters=None):
    """搜索记录（委托给 db 模块，支持多字段过滤）"""
    return db.search_records(query, page, per_page, filters)


# ============ HTML 前端 ============
HTML_PAGE = r"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>基站工参管理器</title>
<style>
*{margin:0;padding:0;box-sizing:border-box}
:root{--primary:#2563eb;--primary-dark:#1d4ed8;--bg:#f1f5f9;--card:#fff;--text:#1e293b;--text-sec:#64748b;--border:#e2e8f0;--success:#16a34a;--danger:#dc2626;--warn:#d97706;--radius:10px}
body{font-family:-apple-system,BlinkMacSystemFont,'Segoe UI','PingFang SC','Microsoft YaHei',sans-serif;background:var(--bg);color:var(--text);min-height:100vh}
.header{background:linear-gradient(135deg,#1e40af,#3b82f6);color:#fff;padding:20px 0;box-shadow:0 2px 10px rgba(0,0,0,.15)}
.header h1{font-size:24px;font-weight:700;letter-spacing:1px}
.header p{font-size:13px;opacity:.85;margin-top:4px}
.container{max-width:1400px;margin:0 auto;padding:20px}
.stats-bar{display:flex;gap:16px;margin-bottom:20px;flex-wrap:wrap}
.stat-card{background:var(--card);border-radius:var(--radius);padding:16px 24px;box-shadow:0 1px 3px rgba(0,0,0,.08);min-width:150px;flex:1}
.stat-card .label{font-size:12px;color:var(--text-sec);text-transform:uppercase;letter-spacing:.5px}
.stat-card .value{font-size:28px;font-weight:700;color:var(--primary);margin-top:4px}
.search-box{background:var(--card);border-radius:var(--radius);padding:20px 24px;box-shadow:0 1px 3px rgba(0,0,0,.08);margin-bottom:20px}
.search-row{display:flex;gap:12px;align-items:center}
.search-input{flex:1;padding:12px 16px;border:2px solid var(--border);border-radius:8px;font-size:15px;outline:none;transition:border-color .2s}
.search-input:focus{border-color:var(--primary)}
.search-grid{display:grid;grid-template-columns:1fr 1fr;gap:10px 16px}
.search-grid label{font-size:12px;color:var(--text-sec);font-weight:600;display:block;margin-bottom:3px}
.search-grid input,.search-grid select{width:100%;padding:10px 12px;border:2px solid var(--border);border-radius:8px;font-size:14px;outline:none;transition:border-color .2s;background:var(--card)}
.search-grid input:focus,.search-grid select:focus{border-color:var(--primary)}
.search-actions{display:flex;gap:10px;margin-top:14px;justify-content:flex-end}
.search-mode-bar{display:flex;justify-content:space-between;align-items:center;margin-bottom:12px}
.mode-label{font-size:14px;font-weight:700;color:var(--primary)}
.btn{padding:10px 20px;border:none;border-radius:8px;font-size:14px;font-weight:600;cursor:pointer;transition:all .2s;display:inline-flex;align-items:center;gap:6px}
.btn-primary{background:var(--primary);color:#fff}.btn-primary:hover{background:var(--primary-dark)}
.btn-success{background:var(--success);color:#fff}.btn-success:hover{opacity:.9}
.btn-danger{background:var(--danger);color:#fff}.btn-danger:hover{opacity:.9}
.btn-outline{background:transparent;border:2px solid var(--border);color:var(--text)}.btn-outline:hover{border-color:var(--primary);color:var(--primary)}
.btn-sm{padding:6px 14px;font-size:12px}
.btn-xs{padding:3px 10px;font-size:11px;border-radius:6px}
.actions-row{display:flex;gap:10px;margin-top:14px;flex-wrap:wrap;align-items:flex-start}
/* 工具栏（导入+文件卡片） */
.toolbar-box{display:flex;align-items:flex-start;gap:12px;margin-bottom:10px;flex-wrap:nowrap}
.toolbar-left{display:flex;gap:8px;flex-shrink:0}
/* 文件管理卡片 */
.file-cards{display:flex;gap:8px;flex-wrap:nowrap;flex:1;min-width:0;overflow-x:auto;padding-bottom:4px}
.file-card{background:#f8fafc;border:1px solid var(--border);border-radius:8px;padding:8px 12px;min-width:180px;max-width:300px;cursor:default;flex-shrink:0}
.file-card-header{display:flex;justify-content:space-between;align-items:center;gap:6px}
.file-card-name{font-size:12px;font-weight:600;color:var(--text);word-break:break-all;flex:1;min-width:0;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
.file-card-meta{font-size:11px;color:var(--text-sec);margin:4px 0;display:flex;align-items:center;gap:6px}
.file-card-sheets{font-size:11px;color:var(--text-sec);cursor:pointer;margin-top:2px;display:flex;align-items:center;gap:4px;user-select:none}
.file-card-sheets:hover{color:var(--primary)}
.file-card-sheets .arrow{font-size:9px;transition:transform .2s;display:inline-block}
.file-card-sheets.expanded .arrow{transform:rotate(90deg)}
.sheet-list{display:none;flex-direction:column;gap:3px;margin-top:6px}
.sheet-list.show{display:flex}
.sheet-item{display:flex;justify-content:space-between;align-items:center;padding:3px 8px;background:#fff;border:1px solid var(--border);border-radius:5px;font-size:11px}
.sheet-item .sheet-name{flex:1;color:var(--text)}
.sheet-item .sheet-count{color:var(--text-sec);font-size:11px;margin-right:8px}
/* 表格 */
.table-wrapper{background:var(--card);border-radius:var(--radius);box-shadow:0 1px 3px rgba(0,0,0,.08);overflow:hidden}
.table-header{display:flex;justify-content:space-between;align-items:center;padding:16px 24px;border-bottom:1px solid var(--border)}
.table-header h3{font-size:16px;font-weight:600}
.table-scroll{overflow-x:auto;max-height:65vh;overflow-y:auto}
table{width:100%;border-collapse:collapse;font-size:13px;table-layout:fixed}
thead{background:#f8fafc;position:sticky;top:0;z-index:2}
th{padding:10px 14px;text-align:left;font-weight:600;color:var(--text-sec);font-size:12px;text-transform:uppercase;letter-spacing:.3px;white-space:nowrap;border-bottom:2px solid var(--border);resize:horizontal;overflow:hidden;min-width:60px}
td{padding:10px 14px;border-bottom:1px solid var(--border);white-space:nowrap;cursor:pointer;user-select:none;overflow:hidden;text-overflow:ellipsis}
td:hover{background:#eef2ff}
tr:hover td{background:#f8fafc}
.carrier-tag{display:inline-block;padding:2px 10px;border-radius:12px;font-size:11px;font-weight:600}
.carrier-dx{background:#dbeafe;color:#1e40af}
.carrier-lt{background:#dcfce7;color:#166534}
.carrier-yd{background:#fef9c3;color:#854d0e}
.tech-tag{display:inline-block;padding:2px 8px;border-radius:4px;font-size:10px;font-weight:700;margin-left:6px}
.tech-5g{background:#7c3aed;color:#fff}
.tech-4g{background:#2563eb;color:#fff}
.vendor-tag{display:inline-block;padding:2px 10px;border-radius:12px;font-size:11px;font-weight:600}
.vendor-hw{background:#fce7f3;color:#9d174d}
.vendor-dt{background:#f3e8ff;color:#6b21a8}
.vendor-ns{background:#e0f2fe;color:#075985}
.vendor-er{background:#fef3c7;color:#92400e}
.vendor-zt{background:#d1fae5;color:#065f46}
.vendor-other{background:#f1f5f9;color:#475569}
.share-tag{display:inline-block;padding:2px 10px;border-radius:12px;font-size:11px;font-weight:600}
.share-yes{background:#fef3c7;color:#92400e}
.share-no{background:#f1f5f9;color:#94a3b8}
.pagination{display:flex;justify-content:center;align-items:center;gap:6px;padding:16px;flex-wrap:wrap}
.page-btn{padding:6px 14px;border:1px solid var(--border);border-radius:6px;background:#fff;cursor:pointer;font-size:13px;transition:all .15s}
.page-btn:hover:not(:disabled){border-color:var(--primary);color:var(--primary)}
.page-btn.active{background:var(--primary);color:#fff;border-color:var(--primary)}
.page-btn:disabled{opacity:.35;cursor:not-allowed}
.page-info{font-size:12px;color:var(--text-sec);margin:0 8px}
.page-jump{width:48px;text-align:center;padding:4px 6px;border:1px solid var(--border);border-radius:6px;font-size:12px}
.empty-state{text-align:center;padding:60px 20px;color:var(--text-sec)}
.empty-state .icon{font-size:48px;margin-bottom:16px;opacity:.3}
.empty-state h3{font-size:18px;margin-bottom:8px;color:var(--text)}
.toast{position:fixed;top:20px;right:20px;padding:12px 20px;border-radius:8px;color:#fff;font-size:14px;font-weight:500;z-index:9999;animation:slideIn .3s ease;box-shadow:0 4px 12px rgba(0,0,0,.15)}
.toast-success{background:var(--success)}.toast-error{background:var(--danger)}.toast-info{background:var(--primary)}
@keyframes slideIn{from{transform:translateX(100%);opacity:0}to{transform:translateX(0);opacity:1}}
.hidden-input{display:none}
/* 外场模式卡片 */
.card-list{display:flex;flex-direction:column;gap:12px;padding:0}
.card-item{background:var(--card);border-radius:var(--radius);padding:14px 16px;box-shadow:0 1px 3px rgba(0,0,0,.08);display:flex;justify-content:space-between;align-items:flex-start;gap:12px;transition:box-shadow .15s}
.card-item:hover{box-shadow:0 2px 8px rgba(0,0,0,.12)}
.card-left{flex:1;min-width:0}
.card-cell-name{font-size:15px;font-weight:700;color:var(--text);line-height:1.4;word-break:break-all}
.card-row{font-size:13px;color:var(--text-sec);margin-top:4px;line-height:1.6}
.card-row span{font-family:monospace;color:var(--text)}
.card-tags{display:flex;gap:6px;margin-bottom:6px}
.card-right{display:flex;flex-direction:column;gap:6px;flex-shrink:0}
.card-btn{padding:5px 12px;border:1px solid var(--border);border-radius:6px;background:var(--card);cursor:pointer;font-size:11px;font-weight:600;color:var(--text-sec);transition:all .15s;white-space:nowrap}
.card-btn:hover{border-color:var(--primary);color:var(--primary)}
.card-btn:active{background:#eff6ff}
.load-more{display:flex;justify-content:center;padding:20px}
.load-more button{padding:10px 32px}
/* 导出下拉菜单 */
.export-dropdown{position:relative;display:inline-flex}
.export-menu{position:absolute;top:100%;right:0;margin-top:4px;background:var(--card);border:1px solid var(--border);border-radius:8px;box-shadow:0 4px 16px rgba(0,0,0,.12);z-index:100;min-width:200px;overflow:hidden}
.export-menu-item{display:block;width:100%;padding:10px 16px;border:none;background:transparent;text-align:left;font-size:13px;cursor:pointer;white-space:nowrap;transition:background .15s}
.export-menu-item:hover{background:#f1f5f9;color:var(--primary)}
/* 遮罩 */
.overlay{position:fixed;top:0;left:0;width:100%;height:100%;background:rgba(0,0,0,.45);z-index:1000;display:flex;align-items:center;justify-content:center}
/* 进度弹窗 */
.progress-dialog{background:var(--card);border-radius:var(--radius);padding:32px 40px;box-shadow:0 8px 32px rgba(0,0,0,.2);min-width:380px;text-align:center}
.spinner{display:inline-block;width:36px;height:36px;border:3px solid var(--border);border-top-color:var(--primary);border-radius:50%;animation:spin .8s linear infinite;margin-bottom:16px}
@keyframes spin{to{transform:rotate(360deg)}}
.progress-dialog .title{font-size:16px;font-weight:600;margin-bottom:8px}
.progress-dialog .detail{font-size:13px;color:var(--text-sec);margin-top:8px;word-break:break-all}
/* 预览弹窗 */
.preview-dialog{background:var(--card);border-radius:var(--radius);padding:0;box-shadow:0 8px 32px rgba(0,0,0,.25);width:620px;max-width:95vw;max-height:85vh;display:flex;flex-direction:column}
.preview-dialog-head{padding:20px 24px 16px;border-bottom:1px solid var(--border);display:flex;justify-content:space-between;align-items:center}
.preview-dialog-head h3{font-size:16px;font-weight:700}
.preview-dialog-body{padding:16px 24px;overflow-y:auto;flex:1}
.preview-dialog-foot{padding:14px 24px;border-top:1px solid var(--border);display:flex;justify-content:flex-end;gap:10px}
.file-section{margin-bottom:20px;border:1px solid var(--border);border-radius:8px;overflow:hidden}
.file-section-head{padding:10px 16px;background:#f8fafc;display:flex;justify-content:space-between;align-items:center;font-size:13px;font-weight:600}
.file-section-meta{font-size:11px;color:var(--text-sec);font-weight:400;margin-left:8px}
.sheet-check-list{padding:8px 16px;display:flex;flex-direction:column;gap:6px}
.sheet-check-item{display:flex;align-items:center;gap:10px;padding:6px 8px;border-radius:6px;cursor:pointer;transition:background .15s}
.sheet-check-item:hover{background:#f1f5f9}
.sheet-check-item input[type=checkbox]{width:16px;height:16px;cursor:pointer;accent-color:var(--primary)}
.sheet-check-item .sc-name{flex:1;font-size:13px}
.sheet-check-item .sc-rows{font-size:11px;color:var(--text-sec);min-width:60px;text-align:right}
.sc-recommended{font-size:10px;background:#dbeafe;color:#1e40af;padding:1px 7px;border-radius:10px;margin-left:6px}
.select-all-row{display:flex;gap:8px;margin-bottom:6px;font-size:12px}
.select-all-row a{color:var(--primary);cursor:pointer;text-decoration:underline}
footer{text-align:center;padding:20px;color:var(--text-sec);font-size:12px}
@media(max-width:768px){.container{padding:12px}.search-grid{grid-template-columns:1fr}.search-row{flex-direction:column}.actions-row{flex-direction:column}}
</style>
</head>
<body>
<div class="header">
  <div class="container" style="display:flex;justify-content:space-between;align-items:center">
    <div><h1>&#x1F4F6; 基站工参管理器</h1><p>快速导入、检索基站工参信息</p></div>
    <div style="font-size:12px;opacity:.7">Station Parameter Manager</div>
  </div>
</div>
<div class="container">
  <div class="stats-bar" id="statsBar">
    <div class="stat-card"><div class="label">已导入小区</div><div class="value" id="totalCount">0</div></div>
    <div class="stat-card"><div class="label">已导入基站</div><div class="value" id="stationCount">0</div></div>
    <div class="stat-card"><div class="label">文件数量</div><div class="value" id="fileCount">0</div></div>
  </div>
  <!-- 文件管理区：导入按钮 + 已载入文件卡片 -->
  <div class="toolbar-box">
    <div class="toolbar-left">
      <button class="btn btn-success" onclick="document.getElementById('fileInput').click()">&#x1F4C1; 导入工参</button>
      <button class="btn btn-outline" onclick="clearAll()">&#x1F5D1; 清空数据</button>
    </div>
    <div class="file-cards" id="fileCards"></div>
  </div>
  <!-- 搜索区：双模式切换 -->
  <div class="search-box">
    <div class="search-mode-bar">
      <span class="mode-label" id="modeLabel">📝 报告模式</span>
      <div style="display:flex;gap:8px;align-items:center">
        <div class="export-dropdown" id="exportDropdown" style="display:none">
          <button class="btn btn-outline btn-xs" onclick="toggleExportMenu()">&#x1F4E5; 导出结果 ▾</button>
          <div class="export-menu" id="exportMenu" style="display:none">
            <button class="export-menu-item" onclick="exportCSV()">📄 表格 CSV（全字段）</button>
            <button class="export-menu-item" onclick="exportPioneer('4G')">📡 Pioneer 4G 基站</button>
            <button class="export-menu-item" onclick="exportPioneer('5G')">📡 Pioneer 5G 基站</button>
            <button class="export-menu-item" onclick="exportAssistant('LTE')">📱 Assistant LTE</button>
            <button class="export-menu-item" onclick="exportAssistant('NR')">📱 Assistant NR</button>
            <button class="export-menu-item" onclick="exportKML('4G')">🗺️ KML 基站扇区 (4G)</button>
            <button class="export-menu-item" onclick="exportKML('5G')">🗺️ KML 基站扇区 (5G)</button>
          </div>
        </div>
        <button class="btn btn-outline btn-xs" onclick="toggleSearchMode()" id="modeToggleBtn">🔀 切换外场模式</button>
      </div>
    </div>
    <div class="search-panel" id="panel_report">
      <div class="search-row">
        <input class="search-input" id="searchInput" type="text" placeholder="输入小区名、PCI、基站ID 或基站名搜索..." autofocus>
        <button class="btn btn-primary" onclick="doSearch()">🔍 搜索</button>
      </div>
    </div>
    <div class="search-panel" id="panel_field" style="display:none">
      <div class="search-grid">
        <div><label>制式</label><select id="f_制式"><option value="">全部</option><option value="4G">4G LTE</option><option value="5G">5G NR</option></select></div>
        <div><label>TAC</label><input id="f_TAC" type="text" placeholder="跟踪区码"></div>
        <div><label>基站ID ⭐</label><input id="f_基站ID" type="text" placeholder="eNodeB / gNodeB ID"></div>
        <div><label>CellID ⭐</label><input id="f_小区ID" type="text" placeholder="小区ID / NR小区标识"></div>
        <div><label>PCI</label><input id="f_PCI" type="text" placeholder="物理小区标识"></div>
        <div><label>频点</label><input id="f_频点" type="text" placeholder="EARFCN / SSB频点"></div>
        <div><label>基站名</label><input id="f_基站名" type="text" placeholder="辅助搜索"></div>
        <div><label>小区名</label><input id="f_小区名" type="text" placeholder="辅助搜索"></div>
      </div>
      <div class="search-actions">
        <button class="btn btn-outline btn-sm" onclick="clearSearch()">✕ 清空</button>
        <button class="btn btn-primary btn-sm" onclick="doSearch(1)">🔍 搜索</button>
      </div>
    </div>
  </div>
  <div class="table-wrapper">
    <div class="table-header">
      <div style="display:flex;align-items:baseline;gap:8px">
        <h3 id="tableTitle" style="white-space:nowrap">&#x1F4CB; 工参数据</h3>
        <span id="searchResultCount" style="font-size:13px;font-weight:400;color:var(--text-sec);white-space:nowrap"></span>
      </div>
    </div>
    <div class="table-scroll">
      <table id="dataTable">
        <thead>
          <tr>
            <th>制式</th><th>运营商</th><th>设备商</th><th>基站名</th><th>基站ID</th>
            <th>小区名</th><th>PCI</th><th>小区ID</th><th>下行频点</th><th>下倾角</th>
            <th>挂高</th><th>方位角</th><th>经度</th><th>纬度</th><th>共享</th><th>TAC</th><th>来源</th>
          </tr>
        </thead>
        <tbody id="tableBody"></tbody>
      </table>
    </div>
    <div id="pagination" class="pagination"></div>
  </div>
  <!-- 外场模式卡片容器 -->
  <div id="cardWrapper" style="display:none">
    <div class="table-header">
      <div style="display:flex;align-items:baseline;gap:8px">
        <h3 id="cardTitle" style="white-space:nowrap">&#x1F4CB; 工参数据</h3>
        <span id="cardResultCount" style="font-size:13px;font-weight:400;color:var(--text-sec);white-space:nowrap"></span>
      </div>
    </div>
    <div id="cardList" class="card-list"></div>
    <div id="cardLoadMore" class="load-more" style="display:none">
      <button class="btn btn-outline btn-sm" onclick="loadMoreCards()">加载更多...</button>
    </div>
  </div>
</div>
<footer>基站工参管理器 &copy; 2026 | 数据仅在本地处理，缓存保存在本地磁盘</footer>

<!-- 进度遮罩 -->
<div class="overlay" id="progressOverlay" style="display:none">
  <div class="progress-dialog">
    <div class="spinner"></div>
    <div class="title" id="progressTitle">正在处理...</div>
    <div class="detail" id="progressDetail">请稍候</div>
  </div>
</div>

<!-- 预览/勾选 弹窗 -->
<div class="overlay" id="previewOverlay" style="display:none">
  <div class="preview-dialog">
    <div class="preview-dialog-head">
      <h3>&#x1F4CB; 选择要导入的工作表</h3>
      <button class="btn btn-outline btn-sm" onclick="closePreview()">✕ 取消</button>
    </div>
    <div class="preview-dialog-body" id="previewBody">加载中...</div>
    <div class="preview-dialog-foot">
      <button class="btn btn-outline" onclick="closePreview()">取消</button>
      <button class="btn btn-success" onclick="confirmImport()">&#x2705; 确认导入</button>
    </div>
  </div>
</div>

<input type="file" id="fileInput" class="hidden-input" accept=".xlsx,.xls" multiple onchange="handleFiles(this.files)">

<script>
let currentPage = 1, lastTotalResults = 0, PER_PAGE = 50;
let searchMode = 'report'; // 'report' 或 'field'
// 预览数据: [{filename, carrier, tech, sheets:[{name,rows,recommended}]}]
let previewData = [];

// ======== 搜索模式切换 ========
function toggleSearchMode() {
  searchMode = (searchMode === 'report') ? 'field' : 'report';
  const panelR = document.getElementById('panel_report');
  const panelF = document.getElementById('panel_field');
  const label = document.getElementById('modeLabel');
  const btn = document.getElementById('modeToggleBtn');
  const tableW = document.querySelector('.table-wrapper');
  const cardW = document.getElementById('cardWrapper');
  if (searchMode === 'field') {
    panelR.style.display = 'none';
    panelF.style.display = 'block';
    label.innerHTML = '🔬 外场模式';
    btn.textContent = '🔀 切换报告模式';
    tableW.style.display = 'none';
    cardW.style.display = 'block';
    document.getElementById('f_基站ID').focus();
    doSearch(1);  // 切换时重新搜索
  } else {
    panelR.style.display = 'block';
    panelF.style.display = 'none';
    label.innerHTML = '📝 报告模式';
    btn.textContent = '🔀 切换外场模式';
    tableW.style.display = 'block';
    cardW.style.display = 'none';
    document.getElementById('searchInput').focus();
  }
}

// 全局键盘事件
document.addEventListener('keydown', e => {
  const tag = document.activeElement.tagName;
  const inInput = (tag === 'INPUT' || tag === 'TEXTAREA' || tag === 'SELECT');
  if (e.key === 'Enter' && inInput) { currentPage = 1; doSearch(); }
  if (e.key === 'Escape' && inInput) {
    clearSearch();
    e.preventDefault();
  }
  // Ctrl+A：聚焦到当前模式的搜索框
  if ((e.ctrlKey || e.metaKey) && e.key === 'a') {
    if (!inInput) {
      e.preventDefault();
      const el = searchMode === 'field' ? document.getElementById('f_基站ID') : document.getElementById('searchInput');
      el.focus(); el.select();
    }
    return;
  }
  // 任意可打印字符：跳转到搜索框
  if (!inInput && !e.ctrlKey && !e.metaKey && !e.altKey && e.key.length === 1) {
    const el = searchMode === 'field' ? document.getElementById('f_基站ID') : document.getElementById('searchInput');
    el.focus();
  }
});

// ======== 统计 & 文件卡片 ========
async function loadStats() {
  try {
    const d = await (await fetch('/api/stats')).json();
    document.getElementById('totalCount').textContent = d.total.toLocaleString();
    document.getElementById('stationCount').textContent = (d.stations || 0).toLocaleString();
    document.getElementById('fileCount').textContent = d.files;
    renderFileCards(d.sources);
  } catch(e) {}
}

function renderFileCards(sources) {
  const fc = document.getElementById('fileCards');
  if (!sources || !sources.length) { fc.innerHTML = ''; return; }
  fc.innerHTML = sources.map(s => {
    const sheets = s.sheets || [];
    const sheetCount = sheets.length;
    const cardId = 'card_' + s.filename.replace(/[^a-zA-Z0-9]/g,'_');
    return `
    <div class="file-card">
      <div class="file-card-header">
        <div class="file-card-name" title="${esc(s.filename)}">${esc(trunc(s.filename, 28))}</div>
        <button class="btn btn-danger btn-xs" onclick="removeFile('${esc(s.filename)}')" title="关闭整个文件">✕</button>
      </div>
      <div class="file-card-meta">
        <span class="carrier-tag carrier-${s.carrier === '中国电信' ? 'dx' : s.carrier === '中国联通' ? 'lt' : 'yd'}">${esc(s.carrier||'')}</span>
        <span class="tech-tag tech-${(s.tech||'').includes('5G')?'5g':'4g'}">${esc(s.tech||'')}</span>
        <span>${s.count.toLocaleString()} 条</span>
      </div>
      ${sheetCount > 0 ? `
      <div class="file-card-sheets" id="${cardId}_toggle" onclick="toggleSheets('${cardId}')">
        <span class="arrow">&#x25B6;</span>
        <span>${sheetCount} 个工作表</span>
      </div>
      <div class="sheet-list" id="${cardId}_sheets">
        ${sheets.map(sn => `
          <div class="sheet-item">
            <span class="sheet-name">${esc(sn)}</span>
            <button class="btn btn-outline btn-xs" style="border-color:#fca5a5;color:#dc2626"
              onclick="removeSheet('${esc(s.filename)}','${esc(sn)}')">移除</button>
          </div>`).join('')}
      </div>` : ''}
    </div>`}).join('');
}

function toggleSheets(cardId) {
  const toggle = document.getElementById(cardId + '_toggle');
  const list = document.getElementById(cardId + '_sheets');
  if (!toggle || !list) return;
  const isOpen = list.classList.contains('show');
  list.classList.toggle('show');
  toggle.classList.toggle('expanded');
}

// ======== 上传 → 预览弹窗 ========
async function handleFiles(files) {
  if (!files.length) return;

  showProgress('正在读取文件...', '解析 sheet 信息，请稍候');

  const fd = new FormData();
  for (let i = 0; i < files.length; i++) fd.append('files', files[i]);
  document.getElementById('fileInput').value = '';

  try {
    const d = await (await fetch('/api/preview', { method: 'POST', body: fd })).json();
    hideProgress();
    if (d.error) { showToast(d.error, 'error'); return; }
    previewData = d.files || [];
    if (!previewData.length) { showToast('未识别到可导入的文件', 'error'); return; }
    openPreview();
  } catch(e) {
    hideProgress();
    showToast('读取文件失败: ' + e.message, 'error');
  }
}

function openPreview() {
  const body = document.getElementById('previewBody');
  if (!previewData.length) return;

  body.innerHTML = previewData.map((f, fi) => {
    if (f.error) return `<div class="file-section">
      <div class="file-section-head" style="color:var(--danger)">${esc(f.filename)} — ${esc(f.error)}</div>
    </div>`;

    const sheetItems = f.sheets.map((s, si) => `
      <label class="sheet-check-item">
        <input type="checkbox" id="chk_${fi}_${si}" data-fi="${fi}" data-si="${si}"
          ${s.recommended ? 'checked' : ''}>
        <span class="sc-name">${esc(s.name)}${s.recommended ? '<span class="sc-recommended">推荐</span>' : ''}</span>
        <span class="sc-rows">${s.rows > 0 ? s.rows.toLocaleString() + ' 行' : ''}</span>
      </label>`).join('');

    return `<div class="file-section">
      <div class="file-section-head">
        ${esc(trunc(f.filename, 40))}
        <span class="file-section-meta">${esc(f.carrier||'')} ${esc(f.tech||'')}</span>
      </div>
      <div class="sheet-check-list">
        <div class="select-all-row">
          <a onclick="selectAll(${fi},true)">全选</a>
          <a onclick="selectAll(${fi},false)">全不选</a>
          <a onclick="selectAll(${fi},'recommended')">仅推荐</a>
        </div>
        ${sheetItems}
      </div>
    </div>`;
  }).join('');

  document.getElementById('previewOverlay').style.display = 'flex';
}

function selectAll(fi, mode) {
  const checks = document.querySelectorAll(`input[data-fi="${fi}"]`);
  checks.forEach(cb => {
    const si = parseInt(cb.getAttribute('data-si'));
    if (mode === 'recommended') cb.checked = previewData[fi].sheets[si].recommended;
    else cb.checked = !!mode;
  });
}

function closePreview() {
  document.getElementById('previewOverlay').style.display = 'none';
  previewData = [];
}

// ======== 确认导入 ========
async function confirmImport() {
  // 收集每个文件选中的 sheet
  const selections = [];
  previewData.forEach((f, fi) => {
    if (f.error) return;
    const selected = [];
    f.sheets.forEach((s, si) => {
      const cb = document.getElementById(`chk_${fi}_${si}`);
      if (cb && cb.checked) selected.push(s.name);
    });
    if (selected.length) selections.push({ filename: f.filename, sheets: selected });
  });

  if (!selections.length) { showToast('请至少勾选一个工作表', 'error'); return; }

  closePreview();
  showProgress('正在导入工参数据...', '准备中');

  // 初始化进度状态
  const totalFiles = selections.length;
  let fileIdx = 0;

  const doImport = async () => {
    for (const sel of selections) {
      fileIdx++;
      document.getElementById('progressDetail').textContent =
        `(${fileIdx}/${totalFiles}) ${sel.filename} — 共 ${sel.sheets.length} 个工作表`;
      try {
        await fetch('/api/import-sheets', {
          method: 'POST',
          headers: { 'Content-Type': 'application/json' },
          body: JSON.stringify(sel)
        });
      } catch(e) {}
    }
  };

  // 启动后台导入，轮询进度
  const r = await fetch('/api/import-sheets-async', {
    method: 'POST',
    headers: { 'Content-Type': 'application/json' },
    body: JSON.stringify({ files: selections })
  });
  const d = await r.json();
  if (d.pending) {
    pollImportProgress();
  } else {
    hideProgress();
    showToast('导入完成', 'success');
    await loadStats(); currentPage = 1; doSearch();
  }
}

// ======== 进度轮询 ========
async function pollImportProgress() {
  try {
    const d = await (await fetch('/api/import-status')).json();
    if (d.total_files > 0) {
      const pct = Math.round(d.current_index / d.total_files * 100);
      document.getElementById('progressTitle').textContent = `正在导入工参数据... (${pct}%)`;
    }
    let msg = `(${d.current_index}/${d.total_files}) ${d.current_file}`;
    if (d.current_sheet) msg += ` → ${d.current_sheet}`;
    msg += ` | 已导入 ${d.imported_records.toLocaleString()} 条`;
    document.getElementById('progressDetail').textContent = msg;
    document.getElementById('totalCount').textContent = d.imported_records.toLocaleString();

    if (d.done || !d.running) {
      hideProgress();
      showToast(`导入完成，共 ${d.imported_records.toLocaleString()} 条记录`, 'success');
      await loadStats(); currentPage = 1; doSearch();
      return;
    }
    setTimeout(pollImportProgress, 500);
  } catch(e) { setTimeout(pollImportProgress, 1000); }
}

// ======== 移除文件 / 移除 sheet ========
async function removeFile(filename) {
  if (!confirm(`确定要移除文件「${filename}」的所有数据吗？`)) return;
  try {
    const r = await fetch('/api/remove-sheet', {
      method: 'POST',
      headers: { 'Content-Type': 'application/json' },
      body: JSON.stringify({ filename })
    });
    const d = await r.json();
    if (d.ok) { showToast('已移除', 'info'); await loadStats(); currentPage = 1; doSearch(); }
  } catch(e) { showToast('操作失败: ' + e.message, 'error'); }
}

async function removeSheet(filename, sheet) {
  if (!confirm(`确定要移除「${filename}」中的工作表「${sheet}」吗？`)) return;
  try {
    const r = await fetch('/api/remove-sheet', {
      method: 'POST',
      headers: { 'Content-Type': 'application/json' },
      body: JSON.stringify({ filename, sheets: [sheet] })
    });
    const d = await r.json();
    if (d.ok) { showToast('已移除工作表', 'info'); await loadStats(); currentPage = 1; doSearch(); }
  } catch(e) { showToast('操作失败: ' + e.message, 'error'); }
}

// ======== 搜索 & 渲染 ========
const FILTER_KEYS = ["制式","TAC","基站ID","小区ID","PCI","频点","基站名","小区名"];

function getFilters() {
  const f = {};
  for (const k of FILTER_KEYS) {
    const el = document.getElementById('f_' + k);
    if (el && el.value.trim()) f[k] = el.value.trim();
  }
  return f;
}

function clearFiltersUI() {
  for (const k of FILTER_KEYS) {
    const el = document.getElementById('f_' + k);
    if (el) el.value = '';
  }
  document.getElementById('f_制式').value = '';
}

// 外场模式卡片数据缓存
let fieldCardPage = 1;
let fieldCardTotal = 0;
let fieldCardResults = [];

function clearSearch() {
  if (searchMode === 'report') {
    document.getElementById('searchInput').value = '';
  } else {
    clearFiltersUI();
  }
  doSearch(1);
}

async function doSearch(page) {
  currentPage = typeof page === 'number' ? page : 1;
  let url = `/api/search?page=${currentPage}&per_page=${PER_PAGE}`;
  let qLabel = '';

  if (searchMode === 'report') {
    const q = document.getElementById('searchInput').value.trim();
    if (q) {
      url += '&q=' + encodeURIComponent(q);
      qLabel = q;
    }
  } else {
    const filters = getFilters();
    for (const [k, v] of Object.entries(filters)) {
      url += '&' + encodeURIComponent(k) + '=' + encodeURIComponent(v);
    }
    if (Object.keys(filters).length > 0) qLabel = '多字段筛选';
  }

  try {
    const resp = await fetch(url);
    if (!resp.ok) throw new Error(`服务器错误 (${resp.status})`);
    const d = await resp.json();
    lastTotalResults = d.total;
    document.getElementById('exportDropdown').style.display = d.total > 0 ? 'inline-flex' : 'none';

    if (searchMode === 'report') {
      document.getElementById('searchResultCount').textContent = d.total > 0 ? `共 ${d.total.toLocaleString()} 条` : '';
      renderTable(d.results, qLabel);
      renderPagination(d.total, d.page, d.pages);
    } else {
      // 外场模式：卡片瀑布流
      document.getElementById('cardResultCount').textContent = d.total > 0 ? `共 ${d.total.toLocaleString()} 条` : '';
      fieldCardPage = d.page;
      fieldCardTotal = d.total;
      fieldCardResults = d.results;
      renderCards(fieldCardResults, fieldCardPage < d.pages);
    }
  } catch(e) {
    console.error('搜索失败:', e);
    showToast('搜索失败: ' + e.message, 'error');
  }
}

function renderCards(data, hasMore) {
  const list = document.getElementById('cardList');
  if (!data || !data.length) {
    list.innerHTML = '<div class="empty-state"><div class="icon">&#x1F50D;</div><h3>未找到匹配记录</h3></div>';
    document.getElementById('cardTitle').innerHTML = '&#x1F4CB; 工参数据';
    document.getElementById('cardLoadMore').style.display = 'none';
    return;
  }
  let html = '';
  for (const r of data) {
    const tech = r['\u6280\u672f\u5236\u5f0f'] || '';
    const techLabel = tech.includes('5G') ? '5G' : '4G';
    const tCls = tech.includes('5G') ? 'tech-5g' : 'tech-4g';
    const cc = r['\u8fd0\u8425\u5546'] || '';
    const shortCarrier = cc.replace('\u4e2d\u56fd', '');
    const cCls = cc.includes('\u7535\u4fe1') ? 'carrier-dx' : cc.includes('\u8054\u901a') ? 'carrier-lt' : 'carrier-yd';
    const cellName = esc(r['\u5c0f\u533a\u540d'] || '');
    const pci = r['PCI'] || '';
    const freq = r['\u4e0b\u884c\u9891\u70b9'] || '';
    const bid = r['\u57fa\u7ad9ID'] || '';
    let cid = r['\u5c0f\u533aID'] || '';
    // 规范化：联通4G 的长格式(基站ID+小区标识)→截取短标识
    if (bid && cid.startsWith(bid)) cid = cid.substring(bid.length);
    // 速查格式：制式_频点_基站ID_Cell ID
    const quickRef = techLabel + '_' + freq + '_' + bid + '_' + cid;

    html += `<div class="card-item">
      <div class="card-left">
        <div class="card-tags">
          <span class="tech-tag ${tCls}">${techLabel}</span>
          <span class="carrier-tag ${cCls}">${shortCarrier}</span>
        </div>
        <div class="card-cell-name">${cellName}</div>
        <div class="card-row">PCI: <span>${pci}</span> &nbsp; 频点: <span>${freq}</span></div>
        <div class="card-row">基站ID: <span>${bid}</span> &nbsp; CellID: <span>${cid}</span></div>
      </div>
      <div class="card-right">
        <button class="card-btn" onclick="copyCellName(this)" data-text="${escAttr(cellName)}">📋 复制小区名</button>
        <button class="card-btn" onclick="copyQuickRef(this)" data-text="${escAttr(quickRef)}">⚡ 复制ID</button>
      </div>
    </div>`;
  }
  list.innerHTML = html;
  document.getElementById('cardTitle').innerHTML = '&#x1F4CB; 工参数据';
  document.getElementById('cardLoadMore').style.display = hasMore ? 'flex' : 'none';
}

async function loadMoreCards() {
  const nextPage = fieldCardPage + 1;
  let url = `/api/search?page=${nextPage}&per_page=${PER_PAGE}`;
  if (searchMode === 'report') return;
  const filters = getFilters();
  for (const [k, v] of Object.entries(filters)) {
    url += '&' + encodeURIComponent(k) + '=' + encodeURIComponent(v);
  }
  try {
    const resp = await fetch(url);
    if (!resp.ok) throw new Error(`服务器错误`);
    const d = await resp.json();
    fieldCardPage = d.page;
    fieldCardTotal = d.total;
    fieldCardResults = fieldCardResults.concat(d.results);
    renderCardsAppend(d.results, fieldCardPage < d.pages);
  } catch(e) {
    showToast('加载失败: ' + e.message, 'error');
  }
}

function renderCardsAppend(data, hasMore) {
  const list = document.getElementById('cardList');
  let html = '';
  for (const r of data) {
    const tech = r['\u6280\u672f\u5236\u5f0f'] || '';
    const techLabel = tech.includes('5G') ? '5G' : '4G';
    const tCls = tech.includes('5G') ? 'tech-5g' : 'tech-4g';
    const cc = r['\u8fd0\u8425\u5546'] || '';
    const shortCarrier = cc.replace('\u4e2d\u56fd', '');
    const cCls = cc.includes('\u7535\u4fe1') ? 'carrier-dx' : cc.includes('\u8054\u901a') ? 'carrier-lt' : 'carrier-yd';
    const cellName = esc(r['\u5c0f\u533a\u540d'] || '');
    const pci = r['PCI'] || '';
    const freq = r['\u4e0b\u884c\u9891\u70b9'] || '';
    const bid = r['\u57fa\u7ad9ID'] || '';
    let cid = r['\u5c0f\u533aID'] || '';
    // 规范化：联通4G 的长格式(基站ID+小区标识)→截取短标识
    if (bid && cid.startsWith(bid)) cid = cid.substring(bid.length);
    const quickRef = techLabel + '_' + freq + '_' + bid + '_' + cid;
    html += `<div class="card-item">
      <div class="card-left">
        <div class="card-tags">
          <span class="tech-tag ${tCls}">${techLabel}</span>
          <span class="carrier-tag ${cCls}">${shortCarrier}</span>
        </div>
        <div class="card-cell-name">${cellName}</div>
        <div class="card-row">PCI: <span>${pci}</span> &nbsp; 频点: <span>${freq}</span></div>
        <div class="card-row">基站ID: <span>${bid}</span> &nbsp; CellID: <span>${cid}</span></div>
      </div>
      <div class="card-right">
        <button class="card-btn" onclick="copyCellName(this)" data-text="${escAttr(cellName)}">📋 复制小区名</button>
        <button class="card-btn" onclick="copyQuickRef(this)" data-text="${escAttr(quickRef)}">⚡ 复制ID</button>
      </div>
    </div>`;
  }
  list.insertAdjacentHTML('beforeend', html);
  document.getElementById('cardLoadMore').style.display = hasMore ? 'flex' : 'none';
}

function copyCellName(btn) {
  const text = btn.getAttribute('data-text');
  navigator.clipboard.writeText(text).then(() => showToast('已复制小区名', 'info'))
    .catch(() => showToast('复制失败', 'error'));
}
function copyQuickRef(btn) {
  const text = btn.getAttribute('data-text');
  navigator.clipboard.writeText(text).then(() => showToast('已复制ID: ' + text, 'info'))
    .catch(() => showToast('复制失败', 'error'));
}
function escAttr(s) { return s.replace(/&/g,'&amp;').replace(/"/g,'&quot;').replace(/</g,'&lt;').replace(/>/g,'&gt;'); }

function renderTable(data, q) {
  const tb = document.getElementById('tableBody');
  if (!data || !data.length) {
    const empty = q ? '未找到匹配的工参记录，请尝试其他关键词' : '暂无数据，请导入工参文件';
    tb.innerHTML = `<tr><td colspan="16"><div class="empty-state"><div class="icon">&#x1F50D;</div><h3>${empty}</h3><p style="margin-top:8px">支持按小区名、PCI、基站ID、基站名搜索</p></div></td></tr>`;
    document.getElementById('tableTitle').innerHTML = '&#x1F4CB; 工参数据';
    return;
  }
  let html = '';
  for (const r of data) {
    const tech = r['\u6280\u672f\u5236\u5f0f'] || '';
    const tCls = tech.includes('5G') ? 'tech-5g' : 'tech-4g';
    const techLabel = tech.includes('5G') ? '5G' : '4G';
    const cc = r['\u8fd0\u8425\u5546'] || '';
    const shortCarrier = cc.replace('\u4e2d\u56fd', '');
    const cCls = cc.includes('\u7535\u4fe1') ? 'carrier-dx' : cc.includes('\u8054\u901a') ? 'carrier-lt' : 'carrier-yd';
    html += `<tr>
      <td><span class="tech-tag ${tCls}">${techLabel}</span></td>
      <td><span class="carrier-tag ${cCls}">${shortCarrier}</span></td>
      <td>${(() => { const v = r['\u8bbe\u5907\u5546']||''; let vc='vendor-other'; if(v.includes('\u534e\u4e3a'))vc='vendor-hw'; else if(v.includes('\u5927\u5510'))vc='vendor-dt'; else if(v.includes('\u8bfa\u57fa\u4e9a')||v.includes('Nokia'))vc='vendor-ns'; else if(v.includes('\u7231\u7acb\u4fe1')||v.includes('Ericsson'))vc='vendor-er'; else if(v.includes('\u4e2d\u5174'))vc='vendor-zt'; return `<span class="vendor-tag ${vc}">${esc(v)}</span>`; })()}</td>
      <td title="${esc(r['\u57fa\u7ad9\u540d']||'')}">${esc(trunc(r['\u57fa\u7ad9\u540d'],30))}</td>
      <td>${esc(r['\u57fa\u7ad9ID']||'')}</td>
      <td title="${esc(r['\u5c0f\u533a\u540d']||'')}">${esc(trunc(r['\u5c0f\u533a\u540d'],35))}</td>
      <td style="font-family:monospace">${esc(r['PCI']||'')}</td>
      <td style="font-family:monospace">${esc(r['\u5c0f\u533aID']||'')}</td>
      <td>${esc(r['\u4e0b\u884c\u9891\u70b9']||'')}</td>
      <td>${esc(r['\u4e0b\u503e\u89d2']||'')}</td>
      <td>${esc(r['\u6302\u9ad8']||'')}</td>
      <td>${esc(r['\u65b9\u4f4d\u89d2']||'')}</td>
      <td>${esc(r['\u7ecf\u5ea6']||'')}</td>
      <td>${esc(r['\u7eac\u5ea6']||'')}</td>
      <td>${(() => { const s = r['\u5171\u4eab']||''; if(!s) return ''; const cls = s.includes('\u975e\u5171\u4eab') ? 'share-no' : 'share-yes'; return `<span class="share-tag ${cls}">${esc(s)}</span>`; })()}</td>
      <td style="font-family:monospace">${esc(r['TAC']||'')}</td>
      <td style="font-size:11px;color:var(--text-sec)">${esc(trunc(r['_\u6587\u4ef6\u540d']||'',25))}</td>
    </tr>`;
  }
  tb.innerHTML = html;
  document.getElementById('tableTitle').innerHTML = '&#x1F4CB; 工参数据' + (q ? ` &mdash; 搜索: <b>${esc(q)}</b> (共 ${lastTotalResults.toLocaleString()} 条)` : '');
}

function renderPagination(total, page, pages) {
  const pg = document.getElementById('pagination');
  if (pages <= 1) { pg.innerHTML = ''; return; }
  let html = '';
  html += `<button class="page-btn" onclick="doSearch(1)"${page<=1?' disabled':''}>&laquo;</button>`;
  html += `<button class="page-btn" onclick="doSearch(${page-1})"${page<=1?' disabled':''}>&lsaquo;</button>`;
  const start = Math.max(1, page-2), end = Math.min(pages, page+2);
  if (start > 1) html += '<span class="page-info">...</span>';
  for (let i = start; i <= end; i++)
    html += `<button class="page-btn${i===page?' active':''}" onclick="doSearch(${i})">${i}</button>`;
  if (end < pages) html += '<span class="page-info">...</span>';
  html += `<button class="page-btn" onclick="doSearch(${page+1})"${page>=pages?' disabled':''}>&rsaquo;</button>`;
  html += `<button class="page-btn" onclick="doSearch(${pages})"${page>=pages?' disabled':''}>&raquo;</button>`;
  html += `<span class="page-info">${page} / ${pages} 页</span>`;
  html += `<span class="page-info">跳至 <input type="number" class="page-jump" id="pageJump" min="1" max="${pages}" placeholder="${page}" onkeydown="if(event.key==='Enter')jumpPage(${pages})"> 页</span>`;
  pg.innerHTML = html;
}

function jumpPage(totalPages) {
  const inp = document.getElementById('pageJump');
  const page = parseInt(inp.value, 10);
  if (!isNaN(page) && page >= 1 && page <= totalPages) {
    doSearch(page);
  } else {
    inp.value = '';
    inp.placeholder = '1~' + totalPages;
  }
}

// ======== 清空 ========
async function clearAll() {
  if (!confirm('确定要清空所有已导入的数据（同时删除本地缓存）？')) return;
  try {
    const d = await (await fetch('/api/clear', { method: 'POST' })).json();
    if (d.ok) {
      showToast('数据已清空', 'info');
      currentPage = 1;
      clearFiltersUI();
      await loadStats(); doSearch();
    }
  } catch(e) { showToast('清空失败: ' + e.message, 'error'); }
}

// ======== 导出 ========
function toggleExportMenu() {
  const menu = document.getElementById('exportMenu');
  menu.style.display = menu.style.display === 'none' ? 'block' : 'none';
}
// 点击页面其他地方关闭菜单
document.addEventListener('click', function(e) {
  const dd = document.getElementById('exportDropdown');
  if (dd && !dd.contains(e.target)) {
    document.getElementById('exportMenu').style.display = 'none';
  }
});

function buildFilterQuery() {
  if (searchMode === 'report') {
    const q = document.getElementById('searchInput').value.trim();
    return q ? 'q=' + encodeURIComponent(q) : '';
  }
  const f = getFilters();
  let qs = '';
  for (const [k, v] of Object.entries(f)) {
    qs += '&' + encodeURIComponent(k) + '=' + encodeURIComponent(v);
  }
  return qs;
}

function exportCSV() {
  document.getElementById('exportMenu').style.display = 'none';
  const a = document.createElement('a');
  a.href = '/api/export?' + buildFilterQuery();
  a.download = '工参查询结果_' + new Date().toISOString().slice(0,10) + '.csv';
  document.body.appendChild(a); a.click(); document.body.removeChild(a);
}

function exportPioneer(fmt) {
  document.getElementById('exportMenu').style.display = 'none';
  const a = document.createElement('a');
  a.href = '/api/export-pioneer?fmt=' + fmt + '&' + buildFilterQuery();
  a.download = 'Pioneer_' + fmt + '_' + new Date().toISOString().slice(0,10) + '.xls';
  document.body.appendChild(a); a.click(); document.body.removeChild(a);
}

function exportAssistant(fmt) {
  document.getElementById('exportMenu').style.display = 'none';
  const a = document.createElement('a');
  a.href = '/api/export-assistant?fmt=' + fmt + '&' + buildFilterQuery();
  a.download = 'Assistant_' + fmt + '_' + new Date().toISOString().slice(0,10) + '.xlsx';
  document.body.appendChild(a); a.click(); document.body.removeChild(a);
}

function exportKML(fmt) {
  document.getElementById('exportMenu').style.display = 'none';
  const a = document.createElement('a');
  a.href = '/api/export-kml?fmt=' + fmt + '&' + buildFilterQuery();
  a.download = (fmt === '5G' ? '5G' : '4G') + '工参_基站扇区_' + new Date().toISOString().slice(0,10) + '.kml';
  document.body.appendChild(a); a.click(); document.body.removeChild(a);
  showToast('正在生成 KML 基站图层...', 'success');
}

// ======== 工具 ========
function showProgress(title, detail) {
  document.getElementById('progressTitle').textContent = title;
  document.getElementById('progressDetail').textContent = detail;
  document.getElementById('progressOverlay').style.display = 'flex';
}
function hideProgress() { document.getElementById('progressOverlay').style.display = 'none'; }
function trunc(s, n) { return s && s.length > n ? s.slice(0,n) + '...' : s || ''; }
function esc(s) {
  if (!s) return '';
  const d = document.createElement('div'); d.textContent = s; return d.innerHTML;
}
function showToast(msg, type) {
  const t = document.createElement('div');
  t.className = 'toast toast-' + type; t.textContent = msg;
  document.body.appendChild(t);
  setTimeout(() => t.remove(), 3000);
}

// ======== 列宽拖拽调整 ========
(function(){
  const table = document.getElementById('dataTable');
  if (!table) return;
  const ths = table.querySelectorAll('thead th');
  ths.forEach((th, i) => {
    const grip = document.createElement('div');
    grip.style.cssText = 'position:absolute;right:0;top:0;bottom:0;width:6px;cursor:col-resize;z-index:3';
    th.style.position = 'relative';
    th.appendChild(grip);
    let startX, startW;
    grip.addEventListener('mousedown', e => {
      e.preventDefault();
      startX = e.clientX;
      startW = th.offsetWidth;
      const onMove = ev => {
        const w = Math.max(60, startW + ev.clientX - startX);
        th.style.width = w + 'px';
        th.style.minWidth = w + 'px';
        th.style.maxWidth = w + 'px';
      };
      const onUp = () => {
        document.removeEventListener('mousemove', onMove);
        document.removeEventListener('mouseup', onUp);
      };
      document.addEventListener('mousemove', onMove);
      document.addEventListener('mouseup', onUp);
    });
  });
})();

// ======== 单元格点击复制 ========
document.getElementById('tableBody').addEventListener('click', function(e) {
  const td = e.target.closest('td');
  if (!td) return;
  // 经纬度列：点击经度或纬度，统一复制"{经度},{纬度}"
  const tr = td.closest('tr');
  const tds = tr ? tr.querySelectorAll('td') : null;
  if (tds && tds.length >= 14) {
    // 经度=第13列(index 12), 纬度=第14列(index 13)
    const lngTd = tds[12], latTd = tds[13];
    if (td === lngTd || td === latTd) {
      const lng = lngTd.innerText.trim();
      const lat = latTd.innerText.trim();
      if (lng && lat) {
        const coord = lng + ',' + lat;
        navigator.clipboard.writeText(coord).then(() => {
          showToast('已复制: ' + coord, 'info');
        }).catch(() => { showToast('复制失败', 'error'); });
        return;
      }
    }
  }
  const text = td.innerText.trim();
  if (!text) return;
  navigator.clipboard.writeText(text).then(() => {
    showToast('已复制: ' + text, 'info');
  }).catch(() => {
    showToast('复制失败', 'error');
  });
});

loadStats(); doSearch();
</script>
</body>
</html>"""


# ============ HTTP 处理器 ============
class Handler(BaseHTTPRequestHandler):
    def log_message(self, format, *args): pass

    def do_GET(self):
        url = urlparse(self.path)
        p = url.path

        if p in ("/", "/index.html"):
            self._send_html(HTML_PAGE)

        elif p == "/api/stats":
            stats = db.get_stats()
            file_sources = db.get_file_sources()
            sources = []
            for fn, v in file_sources.items():
                sources.append({
                    "filename": fn,
                    "carrier": v.get("carrier", ""),
                    "tech": v.get("tech", ""),
                    "count": v.get("count", 0),
                    "sheets": v.get("sheets", []),
                })
            self._send_json(200, {
                "total": stats["total"],
                "stations": stats["stations"],
                "files": stats["files"],
                "sources": sources
            })

        elif p == "/api/import-status":
            st = dict(import_status)
            st["imported_records"] = db.get_total_count()
            self._send_json(200, st)

        elif p == "/api/search":
            params = parse_qs(url.query)
            q = params.get("q", [""])[0]
            page = max(1, int(params.get("page", ["1"])[0]))
            per_page = max(1, int(params.get("per_page", ["50"])[0]))
            filters = _parse_filters(params)
            results, total = search_records(q, page, per_page, filters)
            # 去除 _raw 大字段（前端不需要，避免 JSON 传输膨胀）
            for r in results:
                r.pop("_raw", None)
            pages = max(1, (total + per_page - 1) // per_page)
            self._send_json(200, {"results": results, "total": total, "page": page, "pages": pages})

        elif p == "/api/export":
            params = parse_qs(url.query)
            q = params.get("q", [""])[0]
            filters = _parse_filters(params)
            all_data, _ = search_records(q, 1, 999999, filters)
            cols = ["技术制式", "运营商", "设备商", "基站名", "基站ID", "小区名", "PCI", "小区ID",
                    "下行频点", "下倾角", "挂高", "方位角", "经度", "纬度", "共享", "频段", "TAC"]
            lines = ["\uFEFF" + ",".join(cols)]
            for r in all_data:
                lines.append(",".join('"' + str(r.get(c, "")).replace('"', '""') + '"' for c in cols))
            csv_content = "\n".join(lines)
            self.send_response(200)
            self.send_header("Content-Type", "text/csv; charset=utf-8")
            self.send_header("Content-Disposition",
                             "attachment; filename=export_" + __import__("time").strftime("%Y%m%d_%H%M%S") + ".csv")
            self.end_headers()
            self.wfile.write(csv_content.encode("utf-8"))

        # ===== Pioneer 工参导出（.xls 格式，与 Pioneer 原生模板一致）=====
        elif p == "/api/export-pioneer":
            params = parse_qs(url.query)
            q = params.get("q", [""])[0]
            fmt = params.get("fmt", ["4G"])[0]  # 4G 或 5G
            filters = _parse_filters(params)
            all_data, _ = search_records(q, 1, 999999, filters)

            # 按制式过滤
            tech_keyword = "5G" if fmt == "5G" else "4G"
            filtered = [r for r in all_data if tech_keyword in (r.get("技术制式", "") or "")]

            # 室内外判断
            def is_outdoor(rec):
                name = (rec.get("基站名", "") + rec.get("小区名", "")).lower()
                if any(k in name for k in ("室内", "室分", "indoor")):
                    return "Indoor"
                return "Outdoor"

            if fmt == "5G":
                headers = ["SITE NAME", "CELL NAME", "LONGITUDE", "LATITUDE",
                           "PCI", "SSB ARFCN", "AZIMUTH", "Outdoor/Indoor"]
                rows = [[
                    r.get("基站名", ""),
                    r.get("小区名", ""),
                    r.get("经度", ""),
                    r.get("纬度", ""),
                    r.get("PCI", ""),
                    r.get("下行频点", ""),
                    r.get("方位角", ""),
                    is_outdoor(r),
                ] for r in filtered]
            else:  # 4G
                headers = ["SITE NAME", "CELL NAME", "eNB ID", "LONGITUDE", "LATITUDE",
                           "PCI", "EARFCN", "AZIMUTH", "Outdoor/Indoor"]
                rows = [[
                    r.get("基站名", ""),
                    r.get("小区名", ""),
                    r.get("基站ID", ""),
                    r.get("经度", ""),
                    r.get("纬度", ""),
                    r.get("PCI", ""),
                    r.get("下行频点", ""),
                    r.get("方位角", ""),
                    is_outdoor(r),
                ] for r in filtered]

            # 生成 .xls（xlwt，无编码问题，Pioneer 原生支持）
            import xlwt
            wb = xlwt.Workbook(encoding="utf-8")
            ws = wb.add_sheet(f"{fmt} Site Info", cell_overwrite_ok=True)

            # 写表头
            header_style = xlwt.XFStyle()
            header_font = xlwt.Font()
            header_font.bold = True
            header_style.font = header_font
            for ci, h in enumerate(headers):
                ws.write(0, ci, h, header_style)

            # 写数据行
            for ri, row in enumerate(rows):
                for ci, val in enumerate(row):
                    ws.write(ri + 1, ci, val)

            import io
            buf = io.BytesIO()
            wb.save(buf)
            xls_data = buf.getvalue()

            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.ms-excel")
            self.send_header("Content-Disposition",
                             f"attachment; filename=Pioneer_{fmt}_{ __import__('time').strftime('%Y%m%d_%H%M%S')}.xls")
            self.send_header("Content-Length", str(len(xls_data)))
            self.end_headers()
            self.wfile.write(xls_data)

        # ===== Assistant 工参导出（xlsx 格式）=====
        elif p == "/api/export-assistant":
            params = parse_qs(url.query)
            q = params.get("q", [""])[0]
            fmt = params.get("fmt", ["LTE"])[0]  # LTE 或 NR
            filters = _parse_filters(params)
            all_data, _ = search_records(q, 1, 999999, filters)

            # 按制式过滤
            tech_keyword = "5G" if fmt == "NR" else "4G"
            filtered = [r for r in all_data if tech_keyword in (r.get("技术制式", "") or "")]

            # 室内外判断
            def is_outdoor(rec):
                name = (rec.get("基站名", "") + rec.get("小区名", "")).lower()
                if any(k in name for k in ("室内", "室分", "indoor")):
                    return "Indoor"
                return "Outdoor"

            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active

            if fmt == "NR":
                headers = ["Longitude", "Latitude", "Azimuth", "Height",
                           "Mechanical Downtilt", "gNodeB Name", "gNodeB ID",
                           "Cell Name", "Cell ID", "SsbNarfcn", "PCI", "isOutdoor"]
                ws.title = "NR Site Info"
                ws.append(headers)
                for r in filtered:
                    ws.append([
                        r.get("经度", ""), r.get("纬度", ""), r.get("方位角", ""),
                        r.get("挂高", ""), r.get("下倾角", ""), r.get("基站名", ""),
                        r.get("基站ID", ""), r.get("小区名", ""), r.get("小区ID", ""),
                        r.get("下行频点", ""), r.get("PCI", ""), is_outdoor(r),
                    ])
            else:  # LTE
                headers = ["Longitude", "Latitude", "Azimuth", "eNodeB Name",
                           "eNodeB ID", "Cell Name", "Cell ID", "DlEarfcn",
                           "PCI", "isOutdoor"]
                ws.title = "LTE Site Info"
                ws.append(headers)
                for r in filtered:
                    ws.append([
                        r.get("经度", ""), r.get("纬度", ""), r.get("方位角", ""),
                        r.get("基站名", ""), r.get("基站ID", ""), r.get("小区名", ""),
                        r.get("小区ID", ""), r.get("下行频点", ""), r.get("PCI", ""),
                        is_outdoor(r),
                    ])

            import io
            buf = io.BytesIO()
            wb.save(buf)
            xlsx_data = buf.getvalue()

            self.send_response(200)
            self.send_header("Content-Type",
                             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition",
                             f"attachment; filename=Assistant_{fmt}_{ __import__('time').strftime('%Y%m%d_%H%M%S')}.xlsx")
            self.send_header("Content-Length", str(len(xlsx_data)))
            self.end_headers()
            self.wfile.write(xlsx_data)

        # ===== KML 基站扇区图层导出 =====
        elif p == "/api/export-kml":
            params = parse_qs(url.query)
            q = params.get("q", [""])[0]
            fmt = params.get("fmt", ["4G"])[0]  # 4G 或 5G
            filters = _parse_filters(params)
            all_data, _ = search_records(q, 1, 999999, filters)

            # 按制式过滤
            tech_keyword = "5G" if fmt == "5G" else "4G"
            filtered = [r for r in all_data if tech_keyword in (r.get("技术制式", "") or "")]

            # 生成 KML
            kml_content = kml_export.generate_kml(filtered, fmt)

            import time
            ts = time.strftime("%Y%m%d_%H%M%S")
            filename_ascii = f"{'5G' if fmt == '5G' else '4G'}_Sectors_{ts}.kml"
            # RFC 5987 编码中文文件名
            from urllib.parse import quote
            filename_utf8 = f"{'5G' if fmt == '5G' else '4G'}工参_基站扇区_{ts}.kml"

            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.google-earth.kml+xml; charset=utf-8")
            self.send_header("Content-Disposition",
                             f"attachment; filename=\"{filename_ascii}\"; filename*=UTF-8''{quote(filename_utf8)}")
            self.send_header("Content-Length", str(len(kml_content.encode("utf-8"))))
            self.end_headers()
            self.wfile.write(kml_content.encode("utf-8"))

        else:
            self.send_response(404); self.end_headers()

    def do_POST(self):
        url = urlparse(self.path)
        p = url.path

        if p == "/api/preview":
            self._handle_preview()
        elif p == "/api/import-sheets-async":
            self._handle_import_sheets_async()
        elif p == "/api/remove-sheet":
            self._handle_remove_sheet()
        elif p == "/api/clear":
            db.clear_all()
            # 删除旧缓存文件
            old_cache = os.path.join(WORKSPACE, "cache.json")
            try:
                if os.path.exists(old_cache): os.remove(old_cache)
            except Exception: pass
            self._send_json(200, {"ok": True})
        else:
            self.send_response(404); self.end_headers()

    # ---- 预览：解析 sheet 列表，不导入 ----
    def _handle_preview(self):
        content_type = self.headers.get("Content-Type", "")
        content_length = int(self.headers.get("Content-Length", "0"))
        if not content_length:
            self._send_json(400, {"error": "上传文件为空"}); return

        body = self.rfile.read(content_length)
        boundary = None
        for part in content_type.split(";"):
            part = part.strip()
            if part.startswith("boundary="):
                boundary = part.split("=", 1)[1].strip('"'); break

        if not boundary:
            self._send_json(400, {"error": "无法解析上传数据"}); return

        files_data = self._parse_multipart(body, boundary.encode())
        results = []
        for fi in files_data:
            if not fi["filename"]: continue
            save_path = os.path.join(UPLOAD_DIR, fi["filename"])
            try:
                with open(save_path, "wb") as f:
                    f.write(fi["data"])
                info = preview_file(save_path)
                results.append(info)
            except Exception as e:
                results.append({"filename": fi["filename"], "error": str(e), "sheets": []})

        self._send_json(200, {"files": results})

    # ---- 异步导入指定 sheet ----
    def _handle_import_sheets_async(self):
        body = self._read_json_body()
        if body is None: return
        files = body.get("files", [])  # [{filename, sheets:[...]}, ...]
        if not files:
            self._send_json(400, {"error": "无文件信息"}); return

        import_status.update({
            "running": True, "total_files": len(files),
            "current_index": 0, "current_file": "",
            "current_sheet": "", "done": False, "error": "",
            "imported_records": db.get_total_count(),
        })

        def do_import():
            for i, sel in enumerate(files):
                filename = sel.get("filename", "")
                sheets = sel.get("sheets", [])
                import_status["current_file"] = filename
                import_status["current_index"] = i + 1
                import_status["current_sheet"] = ""
                fp = os.path.join(UPLOAD_DIR, filename)
                if not os.path.exists(fp):
                    import_status["error"] += f"{filename}: 文件不存在; "
                    continue
                try:
                    def _su():
                        def status_update(sheet=None):
                            if sheet: import_status["current_sheet"] = sheet
                        return status_update
                    cnt = import_file_sheets(fp, sheets, status_update=_su())
                    print(f"[导入] {filename}: {cnt} 条")
                except Exception as e:
                    import_status["error"] += f"{filename}: {e}; "

            import_status.update({"running": False, "done": True, "current_sheet": ""})

        threading.Thread(target=do_import, daemon=True).start()
        self._send_json(200, {"pending": True})

    # ---- 移除文件 / sheet ----
    def _handle_remove_sheet(self):
        body = self._read_json_body()
        if body is None: return
        filename = body.get("filename", "")
        sheets = body.get("sheets", None)  # None=删除整个文件
        if not filename:
            self._send_json(400, {"error": "缺少 filename"}); return
        remove_sheets(filename, sheets)
        self._send_json(200, {"ok": True})

    # ---- 工具方法 ----
    def _read_json_body(self):
        try:
            length = int(self.headers.get("Content-Length", "0"))
            raw = self.rfile.read(length)
            return json.loads(raw.decode("utf-8"))
        except Exception as e:
            self._send_json(400, {"error": f"请求体解析失败: {e}"}); return None

    def _parse_multipart(self, body, boundary):
        files = []
        for part in body.split(boundary):
            if len(part) < 10 or b'filename="' not in part: continue
            try:
                header_end = part.find(b"\r\n\r\n")
                if header_end < 0: continue
                header = part[:header_end].decode("utf-8", errors="replace")
                data = part[header_end + 4:]
                if data.endswith(b"\r\n"): data = data[:-2]
                m = re.search(r'filename="([^"]+)"', header)
                if not m or not m.group(1): continue
                files.append({"filename": m.group(1), "data": data})
            except Exception: continue
        return files

    def _send_html(self, content):
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Cache-Control", "no-cache, no-store, must-revalidate")
        self.end_headers()
        self.wfile.write(content.encode("utf-8"))

    def _send_json(self, code, data):
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.end_headers()
        self.wfile.write(json.dumps(data, ensure_ascii=False, default=str).encode("utf-8"))


def open_browser(port):
    import subprocess
    subprocess.Popen(['cmd', '/c', 'start', '', f'http://127.0.0.1:{port}'],
                     creationflags=subprocess.CREATE_NO_WINDOW)


if __name__ == "__main__":
    PORT = 18888

    # 预先绑定端口
    test_sock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    test_sock.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
    test_sock.bind(("127.0.0.1", PORT))
    test_sock.close()

    # 启动时初始化数据库
    init_database()

    server = ThreadingHTTPServer(("127.0.0.1", PORT), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()

    total = db.get_total_count()
    print(f"\n{'='*50}")
    print(f"  基站工参管理器已启动")
    print(f"  访问地址: http://127.0.0.1:{PORT}")
    if total:
        print(f"  已加载 {total} 条记录")
    print(f"{'='*50}\n")

    threading.Timer(0.3, open_browser, args=(PORT,)).start()

    try:
        while True:
            import time; time.sleep(3600)
    except KeyboardInterrupt:
        server.shutdown()
