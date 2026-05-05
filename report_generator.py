"""
道路质量分析优化报告 — 外部生成模块
==============================================
功能：基于工参管理器 API + 问题清单 JSON → 自动生成 docx 报告

用法：
  python report_generator.py <问题清单.json> [--output <输出.docx>] [--api http://127.0.0.1:18888]

依赖：python-docx, requests
"""

import sys
import json
import os
import argparse
import time
from pathlib import Path

try:
    import requests
except ImportError:
    print("[ERROR] 请先安装 requests：pip install requests")
    sys.exit(1)

# 将 office-file-tools 脚本路径加入
SCRIPT_DIR = Path(__file__).parent
SKILL_SCRIPTS = SCRIPT_DIR.parent / ".github" / "skills" / "office-file-tools" / "scripts"
if SKILL_SCRIPTS.exists():
    sys.path.insert(0, str(SKILL_SCRIPTS))

try:
    from create_docx import create_document
except ImportError:
    print("[WARN] 无法导入 create_docx 模块，将使用备用方案")
    create_document = None


# ========== 工参 API 客户端 ==========
class CellParamClient:
    """通过 HTTP API 查询工参管理器"""

    def __init__(self, base_url="http://127.0.0.1:18888"):
        self.base_url = base_url.rstrip("/")

    def query_cell(self, keyword: str, exact=False) -> list[dict]:
        """按小区名模糊查询"""
        try:
            params = {"q": keyword, "limit": "10"}
            if exact:
                params["exact"] = "1"
            resp = requests.get(f"{self.base_url}/api/query-cell", params=params, timeout=10)
            if resp.status_code == 200:
                data = resp.json()
                return data.get("results", [])
            else:
                print(f"  [WARN] API 返回 {resp.status_code}: {resp.text}")
                return []
        except requests.exceptions.ConnectionError:
            print(f"  [ERROR] 无法连接到工参管理器 API ({self.base_url})，请确保 app.py 已启动")
            return []
        except Exception as e:
            print(f"  [WARN] 查询异常: {e}")
            return []

    def find_cell_by_name(self, cell_name: str) -> dict | None:
        """根据小区名查找，支持部分匹配"""
        # 先尝试完整名称匹配
        results = self.query_cell(cell_name)
        for r in results:
            if r["小区名"].lower().replace("-", "") == cell_name.lower().replace("-", ""):
                return r
        # 截取关键段匹配（取 _ 分割的最后两段）
        parts = cell_name.split("_")
        if len(parts) >= 2:
            short = "_".join(parts[-2:])
            results2 = self.query_cell(short)
            for r in results2:
                if short.lower() in r["小区名"].lower():
                    return r
        return results[0] if results else None

    def get_nearby_cells(self, lng: float, lat: float, radius: float = 5.0) -> list[dict]:
        """按经纬度查找附近基站"""
        try:
            params = {"lng": str(lng), "lat": str(lat), "radius": str(radius)}
            resp = requests.get(f"{self.base_url}/api/cell-by-location", params=params, timeout=10)
            if resp.status_code == 200:
                return resp.json().get("results", [])
            return []
        except Exception as e:
            print(f"  [WARN] 附近基站查询异常: {e}")
            return []


# ========== 报告生成引擎 ==========
class ReportGenerator:
    """基于模板和工参数据生成道路质量分析优化报告"""

    def __init__(self, api: CellParamClient):
        self.api = api
        self.missing_cells = []  # 记录未在工参中找到的小区

    def generate(self, problem_data: dict, output_path: str) -> str:
        """
        problem_data 格式：
        {
            "area": "阿乌高速",
            "date": "2026年5月",
            "lte_coverage": "74.04%",
            "lte_avg_rsrp": "-95.82",
            "lte_avg_sinr": "11.23",
            "nr_coverage": "",      // 可选
            "nr_avg_ss_rsrp": "",   // 可选
            "nr_avg_ss_sinr": "",    // 可选
            "test_tools": [
                {"name": "工具名", "purpose": "用途"}
            ],
            "problems": [
                {
                    "type": "覆盖类",          // 覆盖类 | 切换类 | 质差类
                    "network": "4G",           // 4G | 5G
                    "location": "五家渠入口38km处",
                    "cell_name": "CJ_WJQ0_102团8连D_GHCNN_PRT4E_0",
                    "rsrp": "-118.88",
                    "sinr": "",                 // 质差类必填
                    "distance": "8.2km",
                    "root_cause": "服务小区弱覆盖",
                    "cause_detail": "",         // 如 "存在山体阻挡"、"无邻近基站"
                    "nearby_cell": "",          // 邻近小区名
                    "nearby_cell_rsrp": "",     // 邻近小区RSRP
                    "solutions": [
                        "调整{cell}下倾角3°->0°。",
                        "增加{cell}功率。"
                    ]
                }
            ]
        }
        """
        area = problem_data.get("area", "XXX区域")
        date = problem_data.get("date", time.strftime("%Y年%m月"))
        problems = problem_data.get("problems", [])

        print(f"\n{'='*60}")
        print(f"  开始生成报告：{area}道路质量分析优化报告")
        print(f"  问题点数量：{len(problems)}")
        print(f"{'='*60}\n")

        # 第一步：预加载所有涉及的小区工参
        print("[查询] 正在查询工参管理器获取小区参数…")
        cell_cache = {}
        for i, prob in enumerate(problems):
            cell_name = prob.get("cell_name", "")
            if cell_name and cell_name not in cell_cache:
                print(f"  [{i+1}/{len(problems)}] 查询: {cell_name[:60]}…")
                result = self.api.find_cell_by_name(cell_name)
                if result:
                    cell_cache[cell_name] = result
                else:
                    print(f"    [WARN] 未找到！将使用原始名称")
                    self.missing_cells.append(cell_name)
                    cell_cache[cell_name] = None

            nearby = prob.get("nearby_cell", "")
            if nearby and nearby not in cell_cache:
                result = self.api.find_cell_by_name(nearby)
                if result:
                    cell_cache[nearby] = result
                else:
                    cell_cache[nearby] = None

        # 第二步：构建报告内容
        content = self._build_content(area, date, problem_data, problems, cell_cache)
        tables = self._build_tables(problem_data, problems)

        # 第三步：生成 docx
        if create_document:
            create_document(output_path, title=None, content=content, tables=tables)
        else:
            self._fallback_generate(output_path, area, content)

        # 第四步：输出摘要
        self._print_summary(problems)

        return output_path

    def _get_cell_display_name(self, cell_name: str, cache: dict) -> str:
        """获取小区的完整显示名称（工参中的名称优先）"""
        if cell_name in cache and cache[cell_name]:
            return cache[cell_name].get("小区名", cell_name)
        return cell_name

    def _build_content(self, area: str, date: str, data: dict,
                       problems: list, cache: dict) -> list[dict]:
        """构建报告正文内容"""
        content = []

        # === 封面 ===
        content.append({"type": "heading", "level": 0, "text": f"{area}道路质量分析优化报告"})
        content.append({"type": "paragraph", "text": date})
        content.append({"type": "paragraph", "text": "华为技术服务有限公司"})
        content.append({"type": "page_break"})

        # === 一、概述 ===
        content.append({"type": "heading", "level": 1, "text": "一、道路测试优化概述"})
        content.append({"type": "heading", "level": 2, "text": "测试目的"})
        content.append({"type": "paragraph",
                        "text": "通过本轮遍历性测试，利用Assistant平台分析优化，远近协同，"
                                "旨在快速提升目标区域路测指标，减少网内干扰，夯实网络质量，提升用户体验。"})
        content.append({"type": "heading", "level": 2, "text": "测试工具"})
        content.append({"type": "paragraph", "text": "测试工具如下表："})
        content.append({"type": "heading", "level": 2, "text": "测试路线"})
        content.append({"type": "paragraph", "text": f"（{area}道路测试路线图）"})

        # === 二、优化指标 ===
        content.append({"type": "heading", "level": 1, "text": "二、道路优化指标"})
        content.append({"type": "heading", "level": 2, "text": "指标情况"})
        content.append({"type": "paragraph", "text": "本次测试的主要指标如下："})
        content.append({"type": "heading", "level": 2, "text": "测试覆盖图"})
        content.append({"type": "paragraph", "text": f"{area}测试分布图："})
        content.append({"type": "paragraph", "text": "RSRP分布图："})
        content.append({"type": "paragraph", "text": "SINR分布图："})

        # === 三、问题分类汇总 ===
        content.append({"type": "heading", "level": 1, "text": "三、问题原因分类汇总"})
        type_counts = {"覆盖类": 0, "切换类": 0, "质差类": 0}
        for p in problems:
            t = p.get("type", "覆盖类")
            type_counts[t] = type_counts.get(t, 0) + 1
        content.append({"type": "paragraph",
                        "text": f"本次测试共发现 {len(problems)} 个问题点，"
                                f"其中覆盖类 {type_counts['覆盖类']} 个，"
                                f"切换类 {type_counts['切换类']} 个，"
                                f"质差类 {type_counts['质差类']} 个。"})

        # === 四、问题点分析 ===
        content.append({"type": "heading", "level": 1, "text": "四、问题点分析"})

        # 按类型分组
        groups = {"覆盖类": [], "切换类": [], "质差类": []}
        for p in problems:
            t = p.get("type", "覆盖类")
            groups[t].append(p)

        for group_type in ["覆盖类", "切换类", "质差类"]:
            group_problems = groups[group_type]
            if not group_problems:
                continue
            content.append({"type": "heading", "level": 2,
                            "text": f"问题点分析（{group_type}）"})

            for idx, prob in enumerate(group_problems, 1):
                self._add_problem_section(content, prob, idx, cache, area)

        # === 五、总结 ===
        content.append({"type": "heading", "level": 1, "text": "五、测试总结"})
        lte_cov = data.get("lte_coverage", "XX")
        lte_rsrp = data.get("lte_avg_rsrp", "XX")
        lte_sinr = data.get("lte_avg_sinr", "XX")
        summary = (f"本轮测试LTE覆盖率（（-105&-3）*占用时长占比）{lte_cov}，"
                   f"平均RSRP为{lte_rsrp}dBm，平均SINR为{lte_sinr}dB。\n"
                   f"分析问题点{len(problems)}个，"
                   f"其中弱覆盖问题{type_counts['覆盖类']}个，"
                   f"切换问题{type_counts['切换类']}个，"
                   f"质差问题{type_counts['质差类']}个，"
                   f"需推动处理后进行下一轮测试。")
        content.append({"type": "paragraph", "text": summary})

        return content

    def _add_problem_section(self, content: list, prob: dict, idx: int,
                             cache: dict, area: str):
        """添加单个问题点的三段式内容"""
        ptype = prob.get("type", "覆盖类")
        network = prob.get("network", "4G")
        location = prob.get("location", "")
        cell_name = prob.get("cell_name", "")
        rsrp = prob.get("rsrp", "")
        sinr = prob.get("sinr", "")
        distance = prob.get("distance", "")
        root_cause = prob.get("root_cause", "")
        cause_detail = prob.get("cause_detail", "")
        nearby_cell = prob.get("nearby_cell", "")
        nearby_rsrp = prob.get("nearby_cell_rsrp", "")
        solutions = prob.get("solutions", [])

        # 获取工参数据
        cell_info = cache.get(cell_name)
        nearby_info = cache.get(nearby_cell) if nearby_cell else None

        # 小区显示名
        display_cell = self._get_cell_display_name(cell_name, cache)

        # === 标题 ===
        if ptype == "质差类":
            title = f"{area}道路{network}质差：{location}道路{network}质差{idx}"
        elif "/无覆盖" in location or "无覆盖" in root_cause:
            title = f"{area}道路{network}弱覆盖/无覆盖：{location}道路{network}弱覆盖/无覆盖{idx}"
        else:
            title = f"{area}道路{network}弱覆盖：{location}道路{network}弱覆盖{idx}"
        content.append({"type": "heading", "level": 3, "text": title})

        # === 问题描述 ===
        if ptype == "质差类":
            operator_info = ""
            if cell_info:
                operator_info = f"({cell_info.get('运营商', '')})"
            desc = (f"行驶至{location}时，UE占用{operator_info}小区{display_cell}，"
                    f"存在{network}质差，问题路段持续约{distance}。")
        else:
            desc = (f"测试车辆行驶至{location}，UE占用小区{display_cell}时，"
                    f"存在{network}弱覆盖{'/无覆盖' if '无覆盖' in root_cause else ''}问题，"
                    f"问题路段持续约{distance}。")
        content.append({"type": "paragraph", "text": "问题描述："})
        content.append({"type": "paragraph", "text": desc})

        # === 问题分析 ===
        analysis = self._build_analysis(prob, cell_info, nearby_info, cache)
        content.append({"type": "paragraph", "text": "问题分析："})
        content.append({"type": "paragraph", "text": analysis})

        # === 优化方案 ===
        solution_label = "解决方案：" if ptype == "质差类" else "优化方案："
        content.append({"type": "paragraph", "text": solution_label})

        for sol in solutions:
            # 替换模板变量
            sol_text = sol
            sol_text = sol_text.replace("{cell}", display_cell)
            if nearby_cell:
                nearby_display = self._get_cell_display_name(nearby_cell, cache)
                sol_text = sol_text.replace("{nearby_cell}", nearby_display)
            if cell_info:
                sol_text = sol_text.replace("{当前下倾角}", cell_info.get("下倾角", "?"))
                sol_text = sol_text.replace("{当前方位角}", cell_info.get("方位角", "?"))
                sol_text = sol_text.replace("{经度}", cell_info.get("经度", "?"))
                sol_text = sol_text.replace("{纬度}", cell_info.get("纬度", "?"))
            content.append({"type": "number", "text": sol_text})

    def _build_analysis(self, prob: dict, cell_info: dict | None,
                        nearby_info: dict | None, cache: dict) -> str:
        """根据问题类型和工参数据构建分析话术"""
        ptype = prob.get("type", "覆盖类")
        network = prob.get("network", "4G")
        rsrp = prob.get("rsrp", "")
        sinr = prob.get("sinr", "")
        root_cause = prob.get("root_cause", "")
        cause_detail = prob.get("cause_detail", "")
        nearby_cell = prob.get("nearby_cell", "")
        nearby_rsrp = prob.get("nearby_cell_rsrp", "")
        distance = prob.get("distance", "")

        cell_name = prob.get("cell_name", "")
        display_cell = self._get_cell_display_name(cell_name, cache)

        # 附加工参信息
        extra_info = ""
        if cell_info:
            extra_info = (f"（基站:{cell_info.get('基站名','?')}, "
                          f"经度:{cell_info.get('经度','?')}, "
                          f"纬度:{cell_info.get('纬度','?')}, "
                          f"挂高:{cell_info.get('挂高','?')}m, "
                          f"下倾角:{cell_info.get('下倾角','?')}°, "
                          f"方位角:{cell_info.get('方位角','?')}°）")

        if ptype == "覆盖类":
            if cause_detail in ("山体阻挡", "建筑物阻挡", "存在山体阻挡", "存在建筑物阻挡"):
                return (f"UE占用{display_cell}(RSRP:{rsrp})时，"
                        f"该路段存在{cause_detail}，导致此路段{network}弱覆盖。{extra_info}")
            elif nearby_cell and nearby_rsrp:
                nearby_display = self._get_cell_display_name(nearby_cell, cache)
                return (f"UE占用小区{display_cell}时，RSRP{rsrp}，"
                        f"较近小区{nearby_display}对当前道路覆盖较弱，RSRP{nearby_rsrp}，"
                        f"导致此路段{network}弱覆盖。{extra_info}")
            else:
                return (f"UE占用小区{display_cell}，RSRP{rsrp}，"
                        f"{root_cause + '，' if root_cause else ''}"
                        f"导致此路段{network}弱覆盖。{extra_info}")

        elif ptype == "切换类":
            if nearby_cell:
                nearby_display = self._get_cell_display_name(nearby_cell, cache)
                return (f"UE占用小区{display_cell}，未切换至较近且覆盖更好小区{nearby_display}，"
                        f"进而导致弱覆盖。{extra_info}")
            else:
                return (f"UE占用小区{display_cell}，未切换至更优小区，"
                        f"进而导致弱覆盖。{extra_info}")

        elif ptype == "质差类":
            if nearby_cell:
                nearby_display = self._get_cell_display_name(nearby_cell, cache)
                return (f"UE占用{display_cell}(RSRP:{rsrp},SINR:{sinr})时，"
                        f"近点小区{nearby_display}弱覆盖，"
                        f"此路段无主覆盖小区，导致该路段{network}质差。{extra_info}")
            else:
                return (f"UE占用{display_cell}(RSRP:{rsrp},SINR:{sinr})时，"
                        f"此路段无主覆盖小区，导致该路段{network}质差。{extra_info}")

        return root_cause

    def _build_tables(self, data: dict, problems: list) -> list[dict]:
        """构建报告中的表格"""
        tables = []

        # 测试工具表
        test_tools = data.get("test_tools", [])
        if test_tools:
            tables.append({
                "caption": "表1：测试工具清单",
                "headers": ["工具名称", "用途"],
                "rows": [[t.get("name", ""), t.get("purpose", "")] for t in test_tools],
            })
        else:
            tables.append({
                "caption": "表1：测试工具清单",
                "headers": ["工具名称", "用途"],
                "rows": [
                    ["测试手机+Assistant平台", "路测数据采集与分析"],
                    ["工参管理器", "基站工参查询与核验"],
                ],
            })

        # KPI 指标表
        tables.append({
            "caption": "表2：关键指标汇总",
            "headers": ["指标项", "LTE", "NR", "备注"],
            "rows": [
                ["覆盖率(-105&-3)", data.get("lte_coverage", "XX"), data.get("nr_coverage", "—"), ""],
                ["平均RSRP/SS-RSRP", f"{data.get('lte_avg_rsrp', 'XX')}dBm",
                 f"{data.get('nr_avg_ss_rsrp', '—')}dBm" if data.get("nr_avg_ss_rsrp") else "—", ""],
                ["平均SINR/SS-SINR", f"{data.get('lte_avg_sinr', 'XX')}dB",
                 f"{data.get('nr_avg_ss_sinr', '—')}dB" if data.get("nr_avg_ss_sinr") else "—", ""],
                ["问题点总数", "—", "—", f"{len(problems)}个"],
            ],
        })

        return tables

    def _fallback_generate(self, output_path: str, area: str, content: list):
        """备用方案：使用 python-docx 直接生成（当无法导入 create_docx 时）"""
        try:
            from docx import Document
            from docx.shared import Pt, Cm
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc = Document()
            style = doc.styles["Normal"]
            style.font.name = "宋体"
            style.font.size = Pt(11)

            for item in content:
                t = item.get("type", "paragraph")
                text = item.get("text", "")
                if t == "heading":
                    level = min(item.get("level", 1), 9)
                    h = doc.add_heading(text, level=level)
                    if level == 0:
                        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif t == "paragraph":
                    doc.add_paragraph(text)
                elif t == "number":
                    doc.add_paragraph(text, style="List Number")
                elif t == "bullet":
                    doc.add_paragraph(text, style="List Bullet")
                elif t == "page_break":
                    doc.add_page_break()

            doc.save(output_path)
            print(f"[OK] 备用方案已生成: {output_path}")
        except ImportError:
            print("[ERROR] python-docx 未安装，无法生成 docx。请运行: pip install python-docx")
            # 最后兜底：输出纯文本
            txt_path = output_path.replace(".docx", ".txt")
            with open(txt_path, "w", encoding="utf-8") as f:
                for item in content:
                    f.write(item.get("text", "") + "\n\n")
            print(f"[FILE] 已输出纯文本版本: {txt_path}")

    def _print_summary(self, problems: list):
        """输出生成摘要"""
        print(f"\n{'='*60}")
        print("  生成摘要")
        print(f"{'='*60}")

        type_counts = {"覆盖类": 0, "切换类": 0, "质差类": 0}
        for p in problems:
            t = p.get("type", "覆盖类")
            type_counts[t] = type_counts.get(t, 0) + 1

        print(f"  覆盖类: {type_counts['覆盖类']} 个")
        print(f"  切换类: {type_counts['切换类']} 个")
        print(f"  质差类: {type_counts['质差类']} 个")

        if self.missing_cells:
            print(f"\n  [WARN] 以下小区未在工参中找到，已使用原始名称：")
            for c in self.missing_cells:
                print(f"    - {c}")


# ========== 问题清单示例生成 ==========
def generate_example_json(output_path: str):
    """生成问题清单 JSON 模板文件"""
    example = {
        "area": "阿乌高速",
        "date": "2026年5月",
        "lte_coverage": "74.04%",
        "lte_avg_rsrp": "-95.82",
        "lte_avg_sinr": "11.23",
        "nr_coverage": "",
        "nr_avg_ss_rsrp": "",
        "nr_avg_ss_sinr": "",
        "test_tools": [
            {"name": "测试手机 + Assistant平台", "purpose": "路测数据采集与优化分析"},
            {"name": "工参管理器", "purpose": "基站工参查询"},
        ],
        "problems": [
            {
                "type": "覆盖类",
                "network": "4G",
                "location": "五家渠入口38km处",
                "cell_name": "CJ_WJQ0_102团8连D_GHCNN_PRT4E_0",
                "rsrp": "-118.88",
                "sinr": "",
                "distance": "8.2km",
                "root_cause": "服务小区弱覆盖",
                "cause_detail": "",
                "nearby_cell": "",
                "nearby_cell_rsrp": "",
                "solutions": [
                    "调整{cell}下倾角3°->0°。",
                    "增加{cell}功率。",
                ],
            },
            {
                "type": "切换类",
                "network": "4G",
                "location": "五家渠入口50km处",
                "cell_name": "CJWJQ五家渠103团14连L_JYH1B",
                "rsrp": "-108.50",
                "sinr": "",
                "distance": "1.8km",
                "root_cause": "未切换至更优小区",
                "cause_detail": "",
                "nearby_cell": "CJ_WJQ0_五家渠103团13连D_GHCNN_PRT5C_10",
                "nearby_cell_rsrp": "-95.00",
                "solutions": [
                    "添加{cell}到{nearby_cell}双向邻区关系。",
                ],
            },
            {
                "type": "质差类",
                "network": "4G",
                "location": "北海西街食品厂小区附近",
                "cell_name": "CJ_WJQ0_青湖名邸楼顶D_GHBXB_PRB2E_0",
                "rsrp": "-98.12",
                "sinr": "-3",
                "distance": "360m",
                "root_cause": "无主覆盖小区",
                "cause_detail": "",
                "nearby_cell": "CJWJQ五家渠北海公园_FYF1A",
                "nearby_cell_rsrp": "-105.00",
                "solutions": [
                    "调整{nearby_cell}方位角0°->30°。",
                    "调整{nearby_cell}下倾角6°->3°。",
                ],
            },
        ],
    }

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(example, f, ensure_ascii=False, indent=2)
    print(f"[OK] 问题清单模板已生成: {output_path}")
    print(f"   请编辑此文件后运行:")
    print(f"   python report_generator.py {output_path}")


# ========== 命令行入口 ==========
def main():
    parser = argparse.ArgumentParser(
        description="道路质量分析优化报告生成器 — 基于工参管理器 API + 问题清单 JSON",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
示例：
  # 生成问题清单模板
  python report_generator.py --example

  # 基于问题清单生成报告
  python report_generator.py problems.json

  # 指定输出路径和 API 地址
  python report_generator.py problems.json --output 输出报告.docx --api http://127.0.0.1:18888
        """,
    )
    parser.add_argument("input", nargs="?", help="问题清单 JSON 文件路径")
    parser.add_argument("--output", "-o", default="", help="输出 docx 文件路径")
    parser.add_argument("--api", default="http://127.0.0.1:18888", help="工参管理器 API 地址")
    parser.add_argument("--example", action="store_true", help="生成问题清单模板 JSON")
    args = parser.parse_args()

    # 生成示例
    if args.example:
        example_path = args.input or "问题清单模板.json"
        generate_example_json(example_path)
        return

    # 验证输入
    if not args.input:
        parser.print_help()
        print("\n[ERROR] 请指定问题清单 JSON 文件，或使用 --example 生成模板")
        sys.exit(1)

    input_path = Path(args.input)
    if not input_path.exists():
        print(f"[ERROR] 文件不存在: {args.input}")
        print("   使用 --example 生成模板文件")
        sys.exit(1)

    # 加载问题清单
    with open(input_path, "r", encoding="utf-8") as f:
        problem_data = json.load(f)

    # 确定输出路径
    if args.output:
        output_path = args.output
    else:
        area = problem_data.get("area", "report")
        output_path = f"{area}道路质量分析优化报告_{time.strftime('%Y%m%d')}.docx"

    # 创建 API 客户端并生成报告
    api = CellParamClient(args.api)
    generator = ReportGenerator(api)
    result = generator.generate(problem_data, output_path)

    print(f"\n[OK] 报告已生成: {os.path.abspath(result)}")


if __name__ == "__main__":
    main()
